"""Áreas de ponderação (AP) do Censo 2022: a unidade dos resultados da AMOSTRA no Perfil × voto (rodada 49).

Religião, nível de instrução, renda per capita, trabalho, internet, migração e deslocamento são perguntas do
questionário da AMOSTRA: o IBGE não as publica por setor nem por bairro, só por área de ponderação
(`tabelas_xlsx.zip`, "Resultados Gerais da Amostra por Áreas de Ponderação"). São 14.270 áreas em todos os
5.570 municípios (Rio 209, São Paulo 332; mediana de ~10.800 pessoas de 10+), com erro amostral — o IBGE
publica os coeficientes de variação à parte.

Código da AP: 10 dígitos = IBGE do município (7) + 3. Os 7 primeiros são o município, como no CD_BAIRRO: o
filtro por município e a regra de ausência de candidatura em prefeito/vereador valem sem mudança.

Ligações:
- setor → AP: exata, pela "Composição das Áreas de Ponderação" (os 468.097 setores da malha);
- local de votação → AP: a AP do setor que CONTÉM o local (`perfil_local.ligar(..., "contem")`).

Indicadores por AP: os da amostra (aqui) + o universo inteiro (renda, cor, densidade, favela, sexo e idade e
o catálogo de `apuracao.censo`), somando os setores da área — `perfil_local.agregar` com a composição como
ligação. Inferência ECOLÓGICA, como nas outras unidades.
"""

from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import polars as pl

from apuracao import ibge
from apuracao import perfil as pf
from apuracao import perfil_local as pfl

logger = logging.getLogger("apuracao.areas_ponderacao")

FONTE_AMOSTRA = "IBGE — Censo 2022 (amostra, área de ponderação)"
COMPOSICAO_PQ = "ibge_censo2022/ap_composicao.parquet"
AMOSTRA_PQ = "ibge_censo2022/ap_amostra.parquet"


@dataclass(frozen=True)
class ItemAmostra:
    rotulo: str
    tabela: str                                     # nome no ZIP ("Tab4_1.xlsx")
    numerador: tuple[int, ...]                      # colunas (0 = código do município) da `tabela`
    denominador: tuple[str, tuple[int, ...]] | None  # (tabela, colunas); None = valor (ex.: renda em R$)


# cabeçalho esperado das colunas usadas (conferido no parse: o IBGE mudou a tabela → erro claro, não número
# errado). Nas tabelas com dois níveis de cabeçalho a coluna do rótulo de cima é a do "Total".
CABECALHOS: dict[str, dict[int, str]] = {
    "Tab4_1.xlsx": {4: "Total", 5: "Católica Apostólica Romana", 6: "Evangélicas", 7: "Espírita",
                    8: "Umbanda e Candomblé", 11: "Sem religião"},
    "Tab2_2.xlsx": {4: "Total", 5: "Sem instrução e Ensino fundamental incompleto", 8: "Ensino superior completo"},
    "Tab8_3.xlsx": {4: "Pessoas", 7: "Renda média", 10: "Renda mediana"},
    "Tab7_1.xlsx": {4: "Total", 5: "Força de trabalho", 6: "Força de trabalho - ocupada",
                    7: "Força de trabalho - desocupada"},
    "Tab1_8.xlsx": {4: "Total", 5: "Sim"},
    "Tab6_1.xlsx": {4: "Pessoas que já residiram fora do município"},
    "Tab3_7.xlsx": {4: "Total"},
    "Tab10_2.xlsx": {4: "Total", 9: "De 61 a 120 minutos", 10: "De 121 a 240 minutos", 11: "Mais de 240 Minutos"},
}
AMOSTRA: dict[str, ItemAmostra] = {
    "pct_catolicos": ItemAmostra("% de católicos (10 anos ou mais)", "Tab4_1.xlsx", (5,), ("Tab4_1.xlsx", (4,))),
    "pct_evangelicos": ItemAmostra("% de evangélicos (10 anos ou mais)", "Tab4_1.xlsx", (6,), ("Tab4_1.xlsx", (4,))),
    "pct_sem_religiao": ItemAmostra("% sem religião (10 anos ou mais)", "Tab4_1.xlsx", (11,), ("Tab4_1.xlsx", (4,))),
    "pct_espiritas": ItemAmostra("% de espíritas (10 anos ou mais)", "Tab4_1.xlsx", (7,), ("Tab4_1.xlsx", (4,))),
    "pct_matriz_africana": ItemAmostra("% de umbanda e candomblé (10 anos ou mais)", "Tab4_1.xlsx", (8,),
                                       ("Tab4_1.xlsx", (4,))),
    "pct_superior_amostra": ItemAmostra("% com superior completo (25 anos ou mais; Censo)", "Tab2_2.xlsx", (8,),
                                        ("Tab2_2.xlsx", (4,))),
    "pct_sem_instrucao_amostra": ItemAmostra("% sem instrução ou fundamental incompleto (25 anos ou mais; Censo)",
                                             "Tab2_2.xlsx", (5,), ("Tab2_2.xlsx", (4,))),
    "renda_pc_media": ItemAmostra("Renda domiciliar per capita média (R$)", "Tab8_3.xlsx", (7,), None),
    "renda_pc_mediana": ItemAmostra("Renda domiciliar per capita mediana (R$)", "Tab8_3.xlsx", (10,), None),
    "pct_ocupados": ItemAmostra("% de ocupados (14 anos ou mais)", "Tab7_1.xlsx", (6,), ("Tab7_1.xlsx", (4,))),
    "pct_desocupados": ItemAmostra("Taxa de desocupação (% da força de trabalho)", "Tab7_1.xlsx", (7,),
                                   ("Tab7_1.xlsx", (5,))),
    "pct_internet": ItemAmostra("% de moradores com internet em casa", "Tab1_8.xlsx", (5,), ("Tab1_8.xlsx", (4,))),
    "pct_migrantes": ItemAmostra("% que já morou em outro município", "Tab6_1.xlsx", (4,), ("Tab3_7.xlsx", (4,))),
    "pct_trabalho_1h_mais": ItemAmostra("% dos ocupados a mais de 1 h do trabalho", "Tab10_2.xlsx", (9, 10, 11),
                                        ("Tab10_2.xlsx", (4,))),
}


# --------------------------------------------------------------------------
# Leitura (puro: bytes → quadros)
# --------------------------------------------------------------------------
def _valor(c: Any) -> float | None:
    """Célula do IBGE: número; "-" = zero; "X", "..." e vazio = sem dado (sigilo/não se aplica)."""
    if isinstance(c, (int, float)) and not isinstance(c, bool):
        return float(c)
    if isinstance(c, str) and c.strip() == "-":
        return 0.0
    return None


def _eh_area(c: Any) -> bool:
    return (isinstance(c, int) and 10**9 <= c < 10**10) or (isinstance(c, str) and c.isdigit() and len(c) == 10)


def ler_tabela(conteudo: bytes, nome: str) -> tuple[pl.DataFrame, dict[int, float | None]]:
    """Uma tabela `TabN_M.xlsx` → (uma linha por AP: CD_AP, CD_MUN, NM_MUN, NM_AP, c<i> das colunas de
    `CABECALHOS[nome]`; a linha "Brasil" como {coluna: valor}). Confere o cabeçalho."""
    import openpyxl

    esperado = CABECALHOS[nome]
    ws = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True).worksheets[0]
    cab: dict[int, set[str]] = {}
    brasil: dict[int, float | None] = {}
    linhas = []
    for r in ws.iter_rows(values_only=True):
        cel = lambda i, r=r: r[i] if i < len(r) else None  # noqa: E731 — linha mais curta no modo read_only
        if r and _eh_area(cel(2)):
            linhas.append([str(cel(2)), int(str(cel(2))[:7]), cel(1), cel(3), *[_valor(cel(i)) for i in esperado]])
        elif r and cel(0) == "Brasil":
            brasil = {i: _valor(cel(i)) for i in esperado}
        elif r and not linhas and not brasil:  # bloco de cabeçalho (antes dos dados)
            for i, c in enumerate(r):
                if isinstance(c, str):
                    cab.setdefault(i, set()).add(c.strip())
    errados = {i: rot for i, rot in esperado.items() if rot not in cab.get(i, set())}
    if errados:
        raise ValueError(f"{nome}: cabeçalho mudou no IBGE (esperado {errados}, encontrado "
                         f"{ {i: sorted(cab.get(i, set())) for i in errados} })")
    if not linhas:
        raise ValueError(f"{nome}: nenhuma área de ponderação")
    schema = {"CD_AP": pl.String, "CD_MUN": pl.Int64, "NM_MUN": pl.String, "NM_AP": pl.String,
              **{f"c{i}": pl.Float64 for i in esperado}}
    return pl.DataFrame(linhas, schema=schema, orient="row"), brasil


def indicadores_amostra(tabelas: dict[str, pl.DataFrame]) -> pl.DataFrame:
    """CD_AP, CD_MUN, NM_MUN, NM_AP + um indicador por item de `AMOSTRA` (sem I/O). Denominador zero ou sem
    dado → nulo."""
    base = next(iter(tabelas.values())).select("CD_AP", "CD_MUN", "NM_MUN", "NM_AP")
    out = base
    for chave, it in AMOSTRA.items():
        num = tabelas[it.tabela].select("CD_AP", sum(pl.col(f"c{i}") for i in it.numerador).alias("_N"))
        if it.denominador is None:
            col = num.select("CD_AP", pl.col("_N").alias(chave))
        else:
            tab, cols = it.denominador
            den = tabelas[tab].select("CD_AP", sum(pl.col(f"c{i}") for i in cols).alias("_D"))
            col = num.join(den, on="CD_AP", how="left").select(
                "CD_AP", pl.when(pl.col("_D") > 0).then(100 * pl.col("_N") / pl.col("_D")).alias(chave))
        out = out.join(col, on="CD_AP", how="left")
    return out


def conferir_brasil(tabelas: dict[str, pl.DataFrame], brasil: dict[str, dict[int, float | None]]) -> pl.DataFrame:
    """Soma das APs × linha "Brasil" de cada tabela, por coluna de contagem (valores como renda média não
    somam e ficam de fora). DIF_PCT = diferença relativa em %."""
    linhas = []
    for nome, df in tabelas.items():
        valores = {it.numerador for it in AMOSTRA.values() if it.tabela == nome and it.denominador is None}
        for i in CABECALHOS[nome]:
            if any(i in cols for cols in valores):
                continue
            soma, ref = df[f"c{i}"].sum(), brasil.get(nome, {}).get(i)
            linhas.append({"TABELA": nome, "COLUNA": CABECALHOS[nome][i], "SOMA_AREAS": soma, "BRASIL": ref,
                           "DIF_PCT": None if not ref else 100 * (soma - ref) / ref})
    return pl.DataFrame(linhas)


def ler_composicao(conteudo: bytes) -> pl.DataFrame:
    """'Composição das Áreas de Ponderação.xlsx' → CD_SETOR, CD_AP, CD_MUN."""
    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else "" for c in next(it)]
    if cab[:3] != ["Setor", "Código do Município", "Área de Ponderação"]:
        raise ValueError(f"composição das áreas de ponderação: cabeçalho mudou no IBGE ({cab[:4]})")
    linhas = [(str(r[0]), str(r[2]), int(r[1])) for r in it if r and r[0] is not None and r[2] is not None]
    return pl.DataFrame(linhas, schema={"CD_SETOR": pl.String, "CD_AP": pl.String, "CD_MUN": pl.Int64}, orient="row")


# --------------------------------------------------------------------------
# Derivados no cache (geradores de `ibge._geradores`)
# --------------------------------------------------------------------------
def _gravar(df: pl.DataFrame, destino: Path) -> pl.DataFrame:
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_suffix(".parquet.tmp")
    df.write_parquet(tmp)
    tmp.replace(destino)
    return df


def composicao(cache: Path) -> pl.DataFrame:
    """Setor → área de ponderação (nacional). Cache: ibge_censo2022/ap_composicao.parquet."""
    destino = cache / COMPOSICAO_PQ
    if destino.exists():
        return pl.read_parquet(destino)
    xlsx = ibge.caminho("ap_composicao", cache, "BR")
    return _gravar(ler_composicao(xlsx.read_bytes()), destino)


def amostra(cache: Path) -> pl.DataFrame:
    """Indicadores da amostra por área de ponderação (nacional). Confere cada tabela com a linha "Brasil"
    (aviso no log acima de 0,5%). Cache: ibge_censo2022/ap_amostra.parquet (refeito se faltar indicador)."""
    destino = cache / AMOSTRA_PQ
    if destino.exists():
        df = pl.read_parquet(destino)
        if set(AMOSTRA) <= set(df.columns):
            return df
        logger.info("%s sem indicadores novos: refazendo", destino.name)
    zp = ibge.caminho("ap_tabelas", cache, "BR")
    tabelas, brasil = {}, {}
    with zipfile.ZipFile(zp) as z:
        for nome in CABECALHOS:
            tabelas[nome], brasil[nome] = ler_tabela(z.read(nome), nome)
    conf = conferir_brasil(tabelas, brasil)
    for r in conf.filter(pl.col("DIF_PCT").abs() > 0.5).iter_rows(named=True):
        logger.warning("amostra por área: %s/%s soma %.0f × Brasil %.0f (%+.2f%%)", r["TABELA"], r["COLUNA"],
                       r["SOMA_AREAS"], r["BRASIL"], r["DIF_PCT"])
    return _gravar(indicadores_amostra(tabelas), destino)


# --------------------------------------------------------------------------
# Perfil × voto com a área de ponderação como unidade
# --------------------------------------------------------------------------
class PerfilVotoArea(pfl.PerfilVotoLocal):
    """Como `PerfilVotoLocal`, com a ÁREA DE PONDERAÇÃO como unidade: votos e perfil do TSE somados pelos
    locais da área; universo do Censo somado pelos setores da área; amostra do Censo direto da tabela."""

    CENSO = pfl.INDICADORES_CENSO_LOCAL | {k: it.rotulo for k, it in AMOSTRA.items()}
    FONTES_CENSO = pf.PerfilVoto.FONTES_CENSO | {k: FONTE_AMOSTRA for k in AMOSTRA}
    UNIDADE = "area"

    def __init__(self, comp) -> None:
        super().__init__(comp)
        self._local_area: dict[int, pl.DataFrame] = {}
        self._votos_ap: dict[tuple[int, int, int], pl.DataFrame] = {}
        self._perfil_ap: dict[int, pl.DataFrame] = {}
        self._censo_ap: pl.DataFrame | None = None

    def _composicao(self) -> pl.DataFrame:
        return self.b._carregar(("ap_composicao",), lambda: composicao(self.b.cache))

    def _amostra(self) -> pl.DataFrame:
        return self.b._carregar(("ap_amostra",), lambda: amostra(self.b.cache))

    def local_area(self, ano: int) -> pl.DataFrame:
        """UNIDADE (local de votação) → CD_AP: a área do setor que contém o local."""
        with self.b.trava:
            if ano not in self._local_area:
                lig = pfl.ligar(self.b.uf, self.b.cache, self.locais(ano), "contem")
                self._local_area[ano] = (lig.join(self._composicao().select("CD_SETOR", "CD_AP"), on="CD_SETOR",
                                                  how="inner").select("UNIDADE", "CD_AP").unique("UNIDADE"))
            return self._local_area[ano]

    def _vb(self, ano: int, cargo: int, turno: int) -> pl.DataFrame:
        chave = (ano, cargo, turno)
        with self.b.trava:
            if chave not in self._votos_ap:
                por_local = super()._vb(ano, cargo, turno)  # CD_BAIRRO = UNIDADE do local
                self._votos_ap[chave] = (
                    por_local.join(self.local_area(ano).rename({"UNIDADE": "CD_BAIRRO"}), on="CD_BAIRRO", how="inner")
                    .group_by("CD_AP", "NR_VOTAVEL").agg(pl.col("QT_VOTOS").sum(), pl.col("NM_VOTAVEL").first())
                    .rename({"CD_AP": "CD_BAIRRO"}))
            return self._votos_ap[chave]

    def perfil(self, ano: int) -> pl.DataFrame:
        """Indicadores do TSE por área: média dos locais ponderada pelos eleitores (= soma das contagens)."""
        with self.b.trava:
            if ano not in self._perfil_ap:
                por_local = super().perfil(ano)  # CD_BAIRRO = UNIDADE do local, ELEITORES, % por indicador
                ks = list(pf.INDICADORES_TSE)
                self._perfil_ap[ano] = (
                    por_local.join(self.local_area(ano).rename({"UNIDADE": "CD_BAIRRO"}), on="CD_BAIRRO", how="inner")
                    .group_by("CD_AP").agg(pl.col("ELEITORES").sum(),
                                           *[(pl.col(k) * pl.col("ELEITORES")).sum().alias(k) for k in ks])
                    .with_columns([(pl.col(k) / pl.col("ELEITORES")).alias(k) for k in ks])
                    .rename({"CD_AP": "CD_BAIRRO"}))
            return self._perfil_ap[ano]

    def censo_local(self, ano: int) -> pl.DataFrame:
        """Universo (setores da área, ligação exata) + amostra. Não depende do ano da eleição."""
        with self.b.trava:
            if self._censo_ap is None:
                st = self.b._carregar(("setores",), lambda: pfl.setores(self.b.uf, self.b.cache))
                lig = self._composicao().select("CD_SETOR", pl.col("CD_AP").alias("UNIDADE"))
                universo = pfl.agregar(st, lig).rename({"UNIDADE": "CD_BAIRRO"})
                am = self._amostra().drop("CD_MUN", "NM_MUN", "NM_AP").rename({"CD_AP": "CD_BAIRRO"})
                self._censo_ap = universo.join(am, on="CD_BAIRRO", how="left")
            return self._censo_ap

    def participacao(self, ano: int, cargo: int, turno: int, metrica: str) -> pl.DataFrame:
        """CD_BAIRRO (= área), VALOR, NUM, DEN: abstenção ou comparecimento somando os aptos dos locais da área."""
        from apuracao import mapa_locais as ml
        por_local = ml._participacao(self, ano, cargo, turno, metrica)  # CD_BAIRRO = UNIDADE do local
        return (por_local.join(self.local_area(ano).rename({"UNIDADE": "CD_BAIRRO"}), on="CD_BAIRRO", how="inner")
                .group_by("CD_AP").agg(pl.col("NUM").sum(), pl.col("DEN").sum())
                .select(pl.col("CD_AP").alias("CD_BAIRRO"),
                        pl.when(pl.col("DEN") > 0).then(100 * pl.col("NUM") / pl.col("DEN")).alias("VALOR"), "NUM", "DEN"))

    def _nomes(self) -> dict[str, tuple[str, str]]:
        return {a: (n, m) for a, n, m in self._amostra().select("CD_AP", "NM_AP", "NM_MUN").iter_rows()}

    def municipios(self) -> pl.DataFrame:
        """Municípios da UF com áreas que têm local de votação (cadastro mais recente carregado; 2026)."""
        anos = sorted(self._locais) or [2026]
        areas = self.local_area(anos[-1]).select(pl.col("CD_AP").unique())
        return (areas.join(self._amostra().select("CD_AP", "CD_MUN", "NM_MUN"), on="CD_AP", how="inner")
                .group_by(pl.col("CD_MUN").cast(pl.Int64), "NM_MUN").len("BAIRROS").sort("NM_MUN"))


# --------------------------------------------------------------------------
# Mapa por área de ponderação (TODO 25, rodada 50)
# --------------------------------------------------------------------------
MALHA = "malhas/areas_ponderacao_{uf}.geojson"
SIMPLIFICAR_GRAUS = 0.0006  # ~60 m; com 5 casas (~1 m) o RJ fica com ~1,5 MB e SP com ~6 MB (bairros RJ: 2,7 MB)
CAMADAS_MAPA = {"voto": "Voto", "perfil": "Perfil (Censo e eleitorado)", "residuo": "Resíduo do Perfil × voto",
                "variacao": "Variação desde a eleição anterior",  # resíduo e variação: escala divergente (p.p.)
                "transferencia": "Destino dos eliminados (1º → 2º turno)"}


def transferencia_por_area(pva: PerfilVotoArea, ano: int, cargo: int) -> tuple[Any, pl.DataFrame]:
    """A inferência 1º → 2º turno por LOCAL (estrato = município) somada por área: os % observados e o voto
    previsto do 1º finalista são médias ponderadas pelos aptos do turno de cada um (= Σ contagens ÷ Σ aptos). O
    previsto é linear nos % do 1º turno; somado assim, é a previsão da área quando os aptos do local são os mesmos
    nos dois turnos — quase sempre (mesmo cadastro; mudam só seções agregadas/não instaladas). Abstenção extra e
    resíduo saem das somas. O destino dos eliminados é estimado por MUNICÍPIO e a área fica dentro de um só
    município: vale o do município (igual em todos os locais da área)."""
    from apuracao import mapa_locais as ml
    res, por = ml.transferencia_por_local(pva, ano, cargo)
    a = res.unidades.cat2[0]
    media = lambda c, aptos: ((pl.col(c) * pl.col(aptos)).sum() / pl.col(aptos).sum()).alias(c)  # noqa: E731
    destinos = [c for c in por.columns if c.startswith("ELIM_PARA_")]
    area = (por.join(pva.local_area(ano), on="UNIDADE", how="inner")
            .group_by("CD_AP").agg(
                media("ABST_1_PCT", "APTOS_1"), media("ELIMINADOS_1_PCT", "APTOS_1"),
                media("ABST_2_PCT", "APTOS_2"), media(f"{a}_2_PCT", "APTOS_2"), media(f"{a}_2_AJUSTE_PCT", "APTOS_2"),
                pl.col("APTOS_1").sum(), pl.col("APTOS_2").sum(), *[pl.col(c).first() for c in destinos])
            .with_columns((pl.col("ABST_2_PCT") - pl.col("ABST_1_PCT")).alias("ABST_EXTRA_PP"),
                          (pl.col(f"{a}_2_PCT") - pl.col(f"{a}_2_AJUSTE_PCT")).alias("RESIDUO_A_PP"))
            .rename({"CD_AP": "UNIDADE"}))
    return res, area


def malha(uf: str, cache: Path) -> dict[str, Any]:
    """GeoJSON das áreas de ponderação da UF: fusão dos setores da malha do IBGE pela composição.
    Propriedades: CD_AP, NM_AP, NM_MUN, CD_MUN. Cache: malhas/areas_ponderacao_<UF>.geojson."""
    import json

    destino = cache / MALHA.format(uf=uf.upper())
    if destino.exists():
        return json.loads(destino.read_text())
    import geopandas as gpd

    setores_zip = ibge.caminho("malha_setores", cache, uf)
    g = gpd.read_file(f"zip://{setores_zip}", columns=["CD_SETOR"])
    area_de = dict(composicao(cache).select("CD_SETOR", "CD_AP").iter_rows())
    g["CD_AP"] = g["CD_SETOR"].astype(str).map(area_de)
    g = g[g["CD_AP"].notna()].to_crs("EPSG:4326")
    d = g[["CD_AP", "geometry"]].dissolve(by="CD_AP").reset_index()
    import shapely
    d["geometry"] = shapely.set_precision(d.geometry.simplify(SIMPLIFICAR_GRAUS, preserve_topology=True).values, 1e-5)
    nomes = {a: (n, m) for a, n, m in amostra(cache).select("CD_AP", "NM_AP", "NM_MUN").iter_rows()}
    d["NM_AP"] = [nomes.get(a, (f"Área {a[-3:]}", None))[0] for a in d["CD_AP"]]  # sem moradores: sem tabela
    d["NM_MUN"] = [nomes.get(a, (None, None))[1] for a in d["CD_AP"]]
    d["CD_MUN"] = [int(a[:7]) for a in d["CD_AP"]]
    geo = json.loads(d.to_json(drop_id=True))
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_suffix(".geojson.tmp")
    tmp.write_text(json.dumps(geo, ensure_ascii=False, separators=(",", ":")))
    tmp.replace(destino)
    return geo


def mapa(pva: PerfilVotoArea, ano: int, camada: str, cargo: int = 3, turno: int = 1, metrica: str | None = None,
         numero: int | None = None, indicador: str | None = None, municipio: int | None = None,
         min_validos: int = 50, ano_ref: int | None = None) -> dict[str, Any]:
    """Valor por área de ponderação, no formato do mapa por bairro (`itens` por código da área).
    "voto": as métricas do mapa por bairro (votos dos locais da área; abstenção/comparecimento somando os aptos);
    "perfil": um indicador da unidade área (TSE, universo e amostra do Censo);
    "residuo": o resíduo do Perfil × voto por área (p.p.; nº de 2 dígitos em proporcional = partido);
    "variacao": a métrica no `ano` − na referência (`ano_ref`, padrão ano − 4), em p.p., pela mesma regra do
    mapa por local (`mapa_locais.variacao`: partido pela entidade). A área não muda entre eleições (é a do Censo
    2022): os locais de cada ano entram na área do setor que os contém;
    "transferencia": o 1º → 2º turno (`metrica` em `mapa_locais.METRICAS_TRANSFERENCIA`), pela inferência por
    local somada na área (`transferencia_por_area`). ValueError = pedido inválido; TseDataError = microdados
    ausentes."""
    from apuracao import bairros as br
    from apuracao import mapa_locais as ml

    if camada not in CAMADAS_MAPA:
        raise ValueError(f"camada deve ser uma de {sorted(CAMADAS_MAPA)}")
    extra: dict[str, Any] = {}
    if camada == "voto":
        if metrica not in br.METRICAS:
            raise ValueError(f"métrica indisponível por área: {metrica}. Use: {', '.join(br.METRICAS)}")
        cargo_nome = br.CARGOS[cargo].title()
        if metrica in br.PARTICIPACAO:
            df = pva.participacao(ano, cargo, turno, metrica).select("CD_BAIRRO", "VALOR")
            rotulo = f"{br.METRICAS[metrica]} — {cargo_nome} {ano}"
        else:
            vb = pva._vb(ano, cargo, turno)
            df = br.metrica(vb, cargo, metrica, numero)
            rotulo = f"{br.METRICAS[metrica].replace('no bairro', 'na área')} — {cargo_nome} {ano}"
            if numero is not None and metrica in ("pct_candidato", "votos_candidato"):
                nome = vb.filter(pl.col("NR_VOTAVEL") == numero)["NM_VOTAVEL"].head(1).to_list()
                rotulo = f"{br.METRICAS[metrica]} — nº {numero}{' ' + nome[0] if nome else ''} ({cargo_nome} {ano})"
            if metrica == "vencedor":
                extra["categorias"] = br.categorias(vb, cargo)
        tipo = "categorico" if metrica == "vencedor" else "sequencial"
        unidade = "%" if metrica.endswith("_pct") or metrica == "pct_candidato" else ""
    elif camada == "perfil":
        indicadores = pva.indicadores()
        if indicador not in indicadores:
            raise ValueError(f"indicador desconhecido: {indicador}")
        df = pva.indicador(indicador, ano).select("CD_BAIRRO", pl.col("X").alias("VALOR"))
        rotulo, tipo = f"{indicadores[indicador]['rotulo']} — por área de ponderação", "sequencial"
        unidade = "%" if indicador.startswith("pct_") else ""
        extra["fonte_indicador"] = indicadores[indicador]["fonte"]
    elif camada == "variacao":
        ano_ref = ano_ref or ano - 4
        if ano_ref >= ano:
            raise ValueError("a eleição de referência deve ser anterior")
        df, rotulo = ml.variacao(pva, ano, ano_ref, cargo, turno, metrica, numero, municipio)
        tipo, unidade = "divergente", "p.p."
        extra.update(ano_ref=ano_ref, lados=[str(ano_ref), str(ano)], sentido="variacao")
    elif camada == "transferencia":
        res, area = transferencia_por_area(pva, ano, cargo)
        df, rotulo, tipo, unidade, mais = ml.camada_transferencia(res, area, ano, cargo, metrica)
        if municipio is not None:
            df = df.filter(pf.municipio_do_bairro() == municipio)
        extra.update(mais)
    else:  # resíduo
        if numero is None or indicador is None:
            raise ValueError("o resíduo precisa do número (candidato; 2 dígitos = partido) e do indicador")
        d = pva.dispersao(ml.alvo(ano, cargo, turno, numero), indicador, min_validos, False, municipio)
        df = d["pontos"].select("CD_BAIRRO", pl.col("RESIDUO").alias("VALOR"), pl.col("Y").alias("VOTO"),
                                pl.col("X").alias("INDICADOR"))
        est = d["estatistica"]
        rotulo, tipo, unidade = f"Resíduo (p.p.): {d['rotulo_y']} × {d['rotulo_x']}", "divergente", "p.p."
        extra.update(sentido="residuo", rotulo_x=d["rotulo_x"], rotulo_y=d["rotulo_y"], min_validos=min_validos,
                     estatistica={k: est.get(k) for k in ("n", "pearson", "r2", "a", "b", "p")})
    com_local = pva.local_area(ano).select(pl.col("CD_AP").unique().alias("CD_BAIRRO"))
    df = df.join(com_local, on="CD_BAIRRO", how="inner")  # só áreas com local de votação (as da UF)
    if municipio is not None:
        df = df.filter(pf.municipio_do_bairro() == municipio)
        com_local = com_local.filter(pf.municipio_do_bairro() == municipio)
    nomes = pva._nomes()
    extras = [c for c in ("VOTO", "INDICADOR", "ANTES", "DEPOIS") if c in df.columns]  # para a dica do mapa
    itens = {r["CD_BAIRRO"]: {"valor": r["VALOR"], "municipio": " — ".join(nomes.get(r["CD_BAIRRO"], ("?", "?"))),
                              "rotulo": r.get("ROTULO"), **{c.lower(): r[c] for c in extras}}
             for r in df.drop_nulls("VALOR").iter_rows(named=True)}
    return {"camada": camada, "metrica": metrica if camada == "voto" else None, "tipo": tipo, "rotulo": rotulo,
            "unidade": unidade, "ano": ano, "itens": itens,
            "cobertura": {"areas_com_dado": len(itens), "areas": com_local.height}, **extra}
