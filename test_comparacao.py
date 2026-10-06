"""Comparação por município entre duas eleições (2022 importado × coletor 2026) — offline."""

from __future__ import annotations

from pathlib import Path

import polars as pl
import pytest
from fastapi.testclient import TestClient

from apuracao import comparacao as cp
from apuracao import historico as h
from apuracao.divulgacao.cliente import ClienteDivulgacao, LimitadorTaxa
from apuracao.divulgacao.coletor import Coletor
from apuracao.web.app import create_app
from conftest import FakeTSE
from test_historico import NIT, RIO, VOTOS_BR, VOTOS_UF, _cand, _detalhe, _votos

IBGE_RIO, IBGE_NIT = 3304557, 3303302


@pytest.fixture()
def dirs(fake_tse: FakeTSE, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Path, Path]:
    """(2022 importado dos microdados sintéticos, 2026 coletado do TSE falso)."""
    monkeypatch.setattr(h, "load_detalhe", lambda ano, cache: _detalhe())
    monkeypatch.setattr(h, "load_candidatos", lambda ano, cache: _cand())
    monkeypatch.setattr(h, "load_votos", lambda ano, uf, cache: (_votos(VOTOS_UF), _votos(VOTOS_BR)))
    monkeypatch.setattr(h, "fonte_padrao", lambda ano, uf, cache: "secao")  # os votos acima são da fonte "secao"
    monkeypatch.setattr(h, "municipios_tse_ibge", lambda uf, cache: pl.DataFrame(
        {"UF": ["RJ", "RJ"], "CD_MUNICIPIO": [RIO, NIT], "CD_MUNICIPIO_IBGE": [IBGE_RIO, IBGE_NIT],
         "NM_MUNICIPIO": ["RIO DE JANEIRO", "NITERÓI"], "CAPITAL": [True, False], "ZONAS": ["4,5", "71"]}))
    a = tmp_path / "historico_2022_t1"
    h.importar(2022, "RJ", 1, tmp_path, a)
    b = tmp_path / "simulado"
    Coletor(ClienteDivulgacao("simulado", sessao=fake_tse, limitador=LimitadorTaxa(1e9)), b).ciclo()
    return a, b


def test_comparar_totais(dirs: tuple[Path, Path]) -> None:
    a, b = cp.carregar_fonte(dirs[0], 0), cp.carregar_fonte(dirs[1], 2026)
    assert (a.ano, b.ano) == (2022, 2026)
    d = cp.comparar(a, b, 3, "abstencao")
    uf = d.filter(pl.col("ABRANGENCIA") == "uf").row(0, named=True)
    assert uf["VALOR_A"] == pytest.approx(20.0)  # 2022 sintético: 360 de 1800
    assert uf["DIF"] == pytest.approx(uf["VALOR_B"] - uf["VALOR_A"])
    rio = d.filter(pl.col("CD_MUNICIPIO") == RIO).row(0, named=True)
    assert rio["CD_MUNICIPIO_IBGE"] == IBGE_RIO and rio["DIF"] is not None
    quissama = d.filter(pl.col("CD_MUNICIPIO") == 58009).row(0, named=True)  # só existe em 2026 nos dados
    assert quissama["VALOR_A"] is None and quissama["DIF"] is None and quissama["NM_MUNICIPIO"] == "QUISSAMÃ"
    ele = cp.comparar(a, b, 3, "eleitorado").filter(pl.col("ABRANGENCIA") == "uf").row(0, named=True)
    assert ele["DIF"] == pytest.approx(100 * (ele["VALOR_B"] - 1800) / 1800)  # variação %


def test_comparar_partido_e_candidato(dirs: tuple[Path, Path]) -> None:
    a, b = cp.carregar_fonte(dirs[0], 0), cp.carregar_fonte(dirs[1], 2026)
    ps = cp.partidos_disponiveis(a, b, 3)
    assert {"PL", "PSB"} <= set(ps["PARTIDO"]) and not ps.filter(pl.col("PARTIDO") == "PL")["NOS_DOIS"].item()
    pl_rio = cp.comparar(a, b, 3, "partido", partido="PL").filter(pl.col("CD_MUNICIPIO") == RIO).row(0, named=True)
    # sem o partido em 2026 (nem sucessor): "sem dado", não 0% (rodada 41)
    assert pl_rio["VALOR_A"] == pytest.approx(100 * 700 / 1050) and pl_rio["VALOR_B"] is None and pl_rio["DIF"] is None
    cand = cp.comparar(a, b, 3, "candidato", numero_a=22, numero_b=66).filter(pl.col("ABRANGENCIA") == "uf").row(0, named=True)
    assert cand["VALOR_A"] == pytest.approx(60.0) and cand["VALOR_B"] > 0
    with pytest.raises(ValueError):
        cp.comparar(a, b, 3, "partido")


def test_api_comparacao(dirs: tuple[Path, Path], tmp_path: Path) -> None:
    site = TestClient(create_app(dirs[1], "RJ", tmp_path, referencia=dirs[0]))
    info = site.get("/api/comparacao/info").json()
    assert info["disponivel"] and (info["ano_a"], info["ano_b"]) == (2022, 2026) and 3 in info["cargos"]
    d = site.get("/api/comparacao?cargo=3&metrica=abstencao").json()
    assert d["unidade"] == "pp" and d["uf"]["VALOR_A"] == pytest.approx(20.0)
    assert set(d["itens"]) == {str(IBGE_RIO), str(IBGE_NIT), "3304151"}
    assert site.get("/api/comparacao?cargo=3&metrica=partido").status_code == 400
    assert site.get("/api/comparacao/partidos?cargo=7").status_code == 200
    sem_ref = TestClient(create_app(dirs[1], "RJ", tmp_path))
    assert not sem_ref.get("/api/comparacao/info").json()["disponivel"]
    assert sem_ref.get("/api/comparacao?cargo=3&metrica=abstencao").status_code == 404


# --------------------------------------------------------------------------
# Histórico do candidato por município (aba Candidato)
# --------------------------------------------------------------------------
def _fonte(ano: int, totais: list[tuple], cands: list[tuple]) -> cp.Fonte:
    """totais: (cargo, abrangência, município, válidos); cands: (cargo, abrangência, município, nº, nome, votos)."""
    tot = pl.DataFrame([{"CARGO": c, "ABRANGENCIA": a, "UF": "RJ", "CD_MUNICIPIO": m, "VALIDOS": v,
                         "DS_CARGO": {6: "Deputado Federal", 7: "Deputado Estadual"}[c], "PCT_SECOES_TOTALIZADAS": 50.0,
                         "TOTALIZACAO_FINAL": False} for c, a, m, v in totais])
    linhas = []
    for c, a, m, n, nome, votos in cands:
        val = next(v for c2, a2, m2, v in totais if (c2, a2, m2) == (c, a, m))
        linhas.append({"CARGO": c, "ABRANGENCIA": a, "UF": "RJ", "CD_MUNICIPIO": m, "NUMERO": n, "NOME_URNA": nome[:5],
                       "NOME": nome, "PARTIDO": "XX", "FEDERACAO": None, "VOTOS": votos,
                       "PCT_VALIDOS": 100 * votos / val, "SITUACAO": None})
    muns = pl.DataFrame({"CD_MUNICIPIO": [RIO, NIT], "CD_MUNICIPIO_IBGE": [IBGE_RIO, IBGE_NIT],
                         "NM_MUNICIPIO": ["RIO DE JANEIRO", "NITERÓI"]})
    cand = pl.DataFrame(linhas).with_columns(pl.col("CD_MUNICIPIO").cast(pl.Int64))  # só "uf": coluna nula
    return cp.Fonte(ano=ano, totais=tot, candidatos=cand, partidos=pl.DataFrame(), municipios=muns)


TOT = [(c, a, m, 1000) for c in (6, 7) for a, m in (("uf", None), ("mun", RIO), ("mun", NIT))]


@pytest.fixture()
def fontes() -> tuple[cp.Fonte, cp.Fonte]:
    """2026: Fulana é a 1234 de deputado federal (sem voto em Niterói). 2022: foi a 77777 de estadual."""
    atual = _fonte(2026, TOT, [(6, "uf", None, 1234, "Fulana de Tal", 300), (6, "mun", RIO, 1234, "Fulana de Tal", 300),
                               (6, "uf", None, 4321, "Beltrano", 500), (6, "mun", RIO, 4321, "Beltrano", 400),
                               (6, "mun", NIT, 4321, "Beltrano", 100)])
    ref = _fonte(2022, TOT, [(7, "uf", None, 77777, "FULANA  DE TAL", 200), (7, "mun", RIO, 77777, "FULANA  DE TAL", 150),
                             (7, "mun", NIT, 77777, "FULANA  DE TAL", 50), (7, "uf", None, 11111, "OUTRO", 900)])
    return atual, ref


def test_historico_casa_pelo_nome_em_outro_cargo(fontes: tuple[cp.Fonte, cp.Fonte]) -> None:
    h = cp.historico_candidato(*fontes, "RJ", 6, 1234)
    assert h.criterio == "nome completo" and (h.anterior["CARGO"], h.anterior["NUMERO"]) == (7, 77777)
    assert h.sufixos == ("2026", "2022") and any("Deputado Estadual" in n for n in h.notas)
    uf = h.tabela.row(0, named=True)
    assert uf["ABRANGENCIA"] == "uf" and uf["NM_MUNICIPIO"] == "RJ" and uf["POSICAO_2026"] == 2
    assert uf["VAR_VOTOS_PCT"] == pytest.approx(50.0) and uf["VAR_PCT_VALIDOS_PP"] == pytest.approx(10.0)
    nit = h.tabela.filter(pl.col("CD_MUNICIPIO") == NIT).row(0, named=True)
    # sem linha em 2026 (arquivo esparso) = 0 voto, 0% e sem posição; variação −100%
    assert (nit["VOTOS_2026"], nit["PCT_VALIDOS_2026"], nit["POSICAO_2026"]) == (0, 0.0, None)
    assert nit["VOTOS_2022"] == 50 and nit["VAR_VOTOS_PCT"] == pytest.approx(-100.0)
    assert nit["CD_MUNICIPIO_IBGE"] == IBGE_NIT and nit["VAR_PCT_VALIDOS_PP"] == pytest.approx(-5.0)


def test_historico_sem_par_homonimo_e_indicado(fontes: tuple[cp.Fonte, cp.Fonte]) -> None:
    atual, ref = fontes
    h = cp.historico_candidato(atual, ref, "RJ", 6, 4321)
    assert h.criterio == "não concorreu" and h.anterior is None
    assert h.tabela["VOTOS_2022"].null_count() == h.tabela.height and h.tabela["VAR_VOTOS_PCT"].null_count() == 3
    assert cp.historico_candidato(atual, None, "RJ", 6, 1234).criterio == "sem referência"
    # homônimo em dois cargos de 2022, nenhum no cargo atual → ambíguo; indicado à mão resolve
    ref2 = _fonte(2022, TOT, [(7, "uf", None, 77777, "Fulana de Tal", 200), (6, "uf", None, 7070, "FULANA DE TAL", 10),
                              (7, "uf", None, 70707, "fulana de tal", 5)])
    h = cp.historico_candidato(atual, ref2, "RJ", 6, 1234)
    assert h.criterio == "nome completo" and h.anterior["NUMERO"] == 7070  # o mesmo cargo desempata
    ref3 = ref2.__class__(**{**ref2.__dict__, "candidatos": ref2.candidatos.filter(pl.col("CARGO") == 7)})
    h = cp.historico_candidato(atual, ref3, "RJ", 6, 1234)
    assert h.criterio == "ambíguo" and [o["NUMERO"] for o in h.opcoes] == [77777, 70707]
    h = cp.historico_candidato(atual, ref3, "RJ", 6, 1234, cargo_ref=7, numero_ref=70707)
    assert h.criterio == "indicado" and h.anterior["VOTOS"] == 5
    with pytest.raises(ValueError):
        cp.historico_candidato(atual, ref3, "RJ", 6, 1234, numero_ref=70707)  # cargo padrão (6): não existe
    with pytest.raises(LookupError):
        cp.historico_candidato(atual, ref, "RJ", 6, 9999)


def test_planilha_historico(fontes: tuple[cp.Fonte, cp.Fonte]) -> None:
    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(cp.planilha_historico(cp.historico_candidato(*fontes, "RJ", 6, 1234),
                                                                   "RJ", "05/10/2026 21:00")))
    assert wb.sheetnames == ["Sobre", "Por município"]
    ws = wb["Por município"]
    cab = [c.value for c in ws[1]]
    assert cab[:3] == ["Município", "Código IBGE", "Votos 2026"] and "Votos 2022" in cab
    uf = {cab[j]: c.value for j, c in enumerate(ws[2])}
    assert uf["Município"] == "RJ" and uf["Variação dos votos"] == pytest.approx(0.5)
    assert uf["Variação % válidos (p.p.)"] == pytest.approx(10.0)
    sobre = {r[0].value: r[1].value for r in wb["Sobre"].iter_rows(min_row=2)}
    assert "77777" in sobre["Candidato em 2022"] and "mesmo nome civil" in sobre["Identificação"]


def test_api_historico_candidato(dirs: tuple[Path, Path], tmp_path: Path) -> None:
    # o governador 66 de 2026 (TSE falso) passa a ter o nome civil do 22 de 2022 (sintético)
    nome = (pl.read_parquet(dirs[1] / "ultimo" / "candidatos.parquet")
            .filter((pl.col("CARGO") == 3) & (pl.col("NUMERO") == 66))["NOME"][0])
    arq = dirs[0] / "ultimo" / "candidatos.parquet"
    pl.read_parquet(arq).with_columns(pl.when((pl.col("CARGO") == 3) & (pl.col("NUMERO") == 22)).then(pl.lit(nome))
                                      .otherwise(pl.col("NOME")).alias("NOME")).write_parquet(arq)
    site = TestClient(create_app(dirs[1], "RJ", tmp_path, referencia=dirs[0]))
    d = site.get("/api/candidato/historico?cargo=3&numero=66").json()
    assert d["criterio"] == "nome completo" and d["anterior"]["NUMERO"] == 22 and d["sufixos"] == ["2026", "2022"]
    uf = d["linhas"][0]
    assert uf["ABRANGENCIA"] == "uf" and uf["PCT_VALIDOS_2022"] == pytest.approx(60.0)
    rio = next(r for r in d["linhas"] if r["CD_MUNICIPIO"] == RIO)
    assert rio["VAR_PCT_VALIDOS_PP"] == pytest.approx(rio["PCT_VALIDOS_2026"] - rio["PCT_VALIDOS_2022"])
    r = site.get("/api/candidato/historico/planilha?cargo=3&numero=66")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert 'filename="historico_66_RJ_2026x2022.xlsx"' in r.headers["content-disposition"]
    assert site.get("/api/candidato/historico?cargo=3&numero=9999").status_code == 404
    assert site.get("/api/candidato/historico?cargo=3&numero=66&numero_ref=9999").status_code == 400
    sem_ref = TestClient(create_app(dirs[1], "RJ", tmp_path)).get("/api/candidato/historico?cargo=3&numero=66").json()
    assert sem_ref["criterio"] == "sem referência" and sem_ref["anterior"] is None


# --------------------------------------------------------------------------
# Variação de partidos por município (gráfico da aba Comparação)
# --------------------------------------------------------------------------
def _fonte_partidos(ano: int, pct: dict[str, list[float]], validos: list[int]) -> cp.Fonte:
    """Majoritário com 1 candidato por partido; `pct[partido]` = % dos válidos em cada município (1..n)."""
    n = len(validos)
    muns = list(range(1, n + 1))
    tot = pl.DataFrame({"CARGO": 1, "ABRANGENCIA": ["uf"] + ["mun"] * n, "UF": "RJ", "CD_MUNICIPIO": [None] + muns,
                        "VALIDOS": [sum(validos)] + validos}).with_columns(pl.col("CD_MUNICIPIO").cast(pl.Int64))
    linhas = []
    for k, (sigla, ps) in enumerate(pct.items()):
        votos = [round(p * v / 100) for p, v in zip(ps, validos)]
        linhas += [{"CARGO": 1, "ABRANGENCIA": "uf", "CD_MUNICIPIO": None, "PARTIDO": sigla, "NUMERO": k, "VOTOS": sum(votos)}]
        linhas += [{"CARGO": 1, "ABRANGENCIA": "mun", "CD_MUNICIPIO": m, "PARTIDO": sigla, "NUMERO": k, "VOTOS": v}
                   for m, v in zip(muns, votos)]
    cand = pl.DataFrame(linhas).with_columns(pl.col("CD_MUNICIPIO").cast(pl.Int64))
    nomes = pl.DataFrame({"CD_MUNICIPIO": muns, "CD_MUNICIPIO_IBGE": [3300000 + m for m in muns],
                          "NM_MUNICIPIO": [f"M{m}" for m in muns]})
    return cp.Fonte(ano=ano, totais=tot, candidatos=cand, partidos=pl.DataFrame(), municipios=nomes)


BASE_PT = [20.0, 30.0, 40.0, 50.0, 60.0, 35.0]
VALIDOS_V = [10_000, 20_000, 40_000, 10_000, 30_000, 100_000]


def test_variacao_uniforme() -> None:
    """+2 p.p. em todo município: b = 1, média 2, DP 0, IC [2, 2] e leitura de deslocamento uniforme."""
    a = _fonte_partidos(2022, {"PT": BASE_PT, "PL": [100 - x for x in BASE_PT]}, VALIDOS_V)
    b = _fonte_partidos(2026, {"PT": [x + 2 for x in BASE_PT], "PL": [98 - x for x in BASE_PT]}, VALIDOS_V)
    d = cp.variacao_partidos(a, b, 1, ["pt", "PL"])
    pt, pl_ = d["partidos"]
    assert pt["partido"] == "PT" and pt["media"] == pytest.approx(2.0) and pt["dp"] == pytest.approx(0, abs=1e-9)
    assert pt["ic_media"] == pytest.approx([2.0, 2.0]) and pt["reta"]["b"] == pytest.approx(1.0)
    assert pt["uf"]["DIF"] == pytest.approx(2.0) and len(pt["pontos"]) == 6
    assert pl_["media"] == pytest.approx(-2.0)
    assert d["butler"]["de"] == "PT" and d["butler"]["uf"] == pytest.approx(-2.0)  # (−2 − 2) / 2
    assert all(r["VALOR"] == pytest.approx(-2.0) for r in d["butler"]["pontos"])
    assert sum(r["DESTAQUE"] for r in pt["pontos"]) == cp.N_DESTAQUES


def test_variacao_proporcional_e_ponderada() -> None:
    """2026 = 1,5 × 2022 − 10 (com ruído): b > 1 com p pequeno; ponderar muda a média."""
    ruido = [0.3, -0.2, 0.1, -0.3, 0.2, -0.1]
    a = _fonte_partidos(2022, {"PT": BASE_PT}, VALIDOS_V)
    b = _fonte_partidos(2026, {"PT": [1.5 * x - 10 + e for x, e in zip(BASE_PT, ruido)]}, VALIDOS_V)
    d = cp.variacao_partidos(a, b, 1, ["PT"])
    r = d["partidos"][0]["reta"]
    assert r["b"] == pytest.approx(1.5, abs=0.02) and r["p_b1"] < 0.001 and r["ic_b"][0] > 1
    assert "redutos" in d["partidos"][0]["leitura"] and d["butler"] is None and not d["ponderado"]
    pond = cp.variacao_partidos(a, b, 1, ["PT"], ponderar=True)["partidos"][0]
    w = [v / sum(VALIDOS_V) for v in VALIDOS_V]
    assert pond["media"] == pytest.approx(sum(wi * (0.5 * x - 10 + e) for wi, x, e in zip(w, BASE_PT, ruido)), abs=0.01)
    assert cp.variacao_partidos(a, b, 1, ["PT"])["partidos"][0]["ic_media"] == d["partidos"][0]["ic_media"]  # semente


def test_variacao_erros() -> None:
    a = _fonte_partidos(2022, {"PT": BASE_PT}, VALIDOS_V)
    with pytest.raises(ValueError, match="sem votos"):
        cp.variacao_partidos(a, a, 1, ["XYZ"])
    with pytest.raises(ValueError, match="1 a 3"):
        cp.variacao_partidos(a, a, 1, ["A", "B", "C", "D"])
    with pytest.raises(ValueError, match="1 a 3"):
        cp.variacao_partidos(a, a, 1, [" "])


def _gravar(f: cp.Fonte, destino: Path) -> Path:
    (destino / "ultimo").mkdir(parents=True)
    for nome in ("totais", "candidatos", "municipios"):
        getattr(f, nome).write_parquet(destino / "ultimo" / f"{nome}.parquet")
    (destino / "status.json").write_text(f'{{"ano": {f.ano}}}')
    return destino


def test_api_variacao(dirs: tuple[Path, Path], tmp_path: Path) -> None:
    a = _gravar(_fonte_partidos(2022, {"PT": BASE_PT, "PL": [100 - x for x in BASE_PT]}, VALIDOS_V), tmp_path / "a")
    b = _gravar(_fonte_partidos(2026, {"PT": [x - 1 for x in BASE_PT], "PL": [101 - x for x in BASE_PT]}, VALIDOS_V),
                tmp_path / "b")
    site = TestClient(create_app(b, "RJ", tmp_path, referencia=a))
    d = site.get("/api/comparacao/variacao?cargo=1&partidos=PT,PL&ponderar=true").json()
    assert (d["ano_a"], d["ano_b"], d["ponderado"]) == (2022, 2026, True)
    assert [p["partido"] for p in d["partidos"]] == ["PT", "PL"] and d["butler"]["uf"] == pytest.approx(1.0)
    assert d["partidos"][0]["pontos"][0]["CD_MUNICIPIO_IBGE"] == 3300001
    assert site.get("/api/comparacao/variacao?cargo=1&partidos=XYZ").status_code == 400
    # 2022 sintético × TSE falso: o PL não tem correspondente em 2026 → 400 com o motivo (rodada 41)
    real = TestClient(create_app(dirs[1], "RJ", tmp_path, referencia=dirs[0]))
    r = real.get("/api/comparacao/variacao?cargo=3&partidos=PL")
    assert r.status_code == 400 and "sem correspondente em 2026" in r.json()["detail"]
    assert TestClient(create_app(dirs[1], "RJ", tmp_path)).get(
        "/api/comparacao/variacao?cargo=3&partidos=PL").status_code == 404



def test_partido_entre_anos_pela_entidade() -> None:
    """Rodada 41: "PC do B" (2022) = "PCDOB" (2026); PRD (2026) = PTB + PATRIOTA (2022); MISSÃO (14, 2026) não é o
    PTB (14, 2022); e a sigla com minúscula funciona no gráfico de variação (antes, .upper() a quebrava)."""
    def fonte(ano: int, pct: dict[tuple[int, str], list[float]]) -> cp.Fonte:
        f = _fonte_partidos(ano, {s_: v for (_, s_), v in pct.items()}, VALIDOS_V)
        nr = {s_: n for n, s_ in pct}
        cand = f.candidatos.with_columns(pl.col("PARTIDO").replace_strict(nr, return_dtype=pl.Int64).alias("NR_PARTIDO"))
        return cp.Fonte(ano=ano, totais=f.totais, candidatos=cand, partidos=f.partidos, municipios=f.municipios)
    a = fonte(2022, {(65, "PC do B"): [10.0] * 6, (14, "PTB"): [5.0] * 6, (51, "PATRIOTA"): [3.0] * 6})
    b = fonte(2026, {(65, "PCDOB"): [12.0] * 6, (25, "PRD"): [6.0] * 6, (14, "MISSÃO"): [4.0] * 6})
    assert cp.siglas_do_partido(a, b, "PCDOB") == (["PC do B"], ["PCDOB"])
    assert cp.siglas_do_partido(a, b, "PRD") == (["PTB", "PATRIOTA"], ["PRD"])
    assert cp.siglas_do_partido(a, b, "PTB") == (["PTB", "PATRIOTA"], ["PRD"])   # pela sigla antiga também
    assert cp.siglas_do_partido(a, b, "MISSÃO") == ([], ["MISSÃO"])
    ps = {r["PARTIDO"]: r for r in cp.partidos_disponiveis(a, b, 1).iter_rows(named=True)}
    assert ps["PRD"]["SIGLAS_A"] == "PTB + PATRIOTA" and ps["PRD"]["NOS_DOIS"] and not ps["MISSÃO"]["NOS_DOIS"]
    uf = cp.comparar(a, b, 1, "partido", partido="PRD").filter(pl.col("ABRANGENCIA") == "uf").row(0, named=True)
    assert uf["VALOR_A"] == pytest.approx(8.0, abs=0.01) and uf["VALOR_B"] == pytest.approx(6.0, abs=0.01)
    d = cp.variacao_partidos(a, b, 1, ["PC do B"])["partidos"][0]
    assert (d["siglas_a"], d["siglas_b"], d["media"]) == (["PC do B"], ["PCDOB"], pytest.approx(2.0, abs=0.01))
    with pytest.raises(ValueError, match="sem correspondente"):
        cp.variacao_partidos(a, b, 1, ["MISSÃO"])


def test_percentual_sobre_os_validos_oficiais_nas_duas_fontes(fontes: tuple[cp.Fonte, cp.Fonte]) -> None:
    """Rodada 42 (TODO 23): o tempo real grava o % do TSE (válidos + sub judice no denominador); comparar fontes
    usa SEMPRE votos ÷ válidos oficiais, recalculado."""
    atual, ref = fontes
    tempo_real = cp.Fonte(atual.ano, atual.totais, atual.candidatos.with_columns(
        (100 * pl.col("VOTOS") / 1100).alias("PCT_VALIDOS")), atual.partidos, atual.municipios)  # 100 sub judice
    uf = cp.comparar(ref, tempo_real, 6, "candidato", numero_a=77777, numero_b=1234).filter(
        pl.col("ABRANGENCIA") == "uf").row(0, named=True)
    assert uf["VALOR_B"] == pytest.approx(30.0)            # 300 / 1000, não 300 / 1100
    h = cp.historico_candidato(tempo_real, ref, "RJ", 6, 1234)
    assert h.atual["PCT_VALIDOS"] == pytest.approx(30.0) and h.tabela.row(0, named=True)["PCT_VALIDOS_2026"] == pytest.approx(30.0)
