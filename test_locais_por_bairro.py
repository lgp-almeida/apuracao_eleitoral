"""Locais e seções de um bairro (`apuracao/locais_bairro.py` + `locais_por_bairro.py`), no cadastro sintético."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import polars as pl
import pytest

import locais_por_bairro as cli
import votos_por_local_votacao as v
from apuracao import bairros as br
from apuracao import eleitorado as el
from apuracao import locais_bairro as lb
from apuracao import versoes as vs
from conftest import escrever_bairros

ZIP_2026 = v.electorate_spec(2026, "RJ").zip_name
CARIMBO = "2026-09-29T062809"  # DT_GERACAO/HH_GERACAO do cadastro sintético


@pytest.fixture
def secoes_2026(tse_cache: Path) -> pl.DataFrame:
    return el.load_sections(2026, "RJ", tse_cache)


def test_cadastro_ignora_acento_e_caixa_e_inclui_agregada(secoes_2026: pl.DataFrame) -> None:
    locais = lb.locais_do_bairro(secoes_2026, "centro")
    # Pedro II (Rio, "Centro"), Escola X (Rio, "CENTRO") e o Pedro II de Niterói: o bairro-texto se repete entre municípios
    assert locais.select("NM_MUNICIPIO", "NR_LOCAL_VOTACAO").rows() == [
        ("NITERÓI", 1015), ("RIO DE JANEIRO", 1015), ("RIO DE JANEIRO", 1023)]
    tab = lb.secoes_dos_locais(secoes_2026, locais.filter(pl.col("NM_MUNICIPIO") == "RIO DE JANEIRO"))
    assert tab.select("NR_LOCAL_VOTACAO", "NR_SECAO", "TIPO_SECAO", "NR_SECAO_PRINCIPAL").rows() == [
        (1015, 10, "PRINCIPAL", None), (1015, 11, "PRINCIPAL", None), (1015, 12, "PRINCIPAL", None),
        (1015, 13, "AGREGADA", 12), (1023, 20, "PRINCIPAL", None)]
    assert tab["QT_ELEITOR_SECAO"].sum() == 310 + 290 + 260 + 95 + 150


def test_exato_por_padrao_e_contem_sob_pedido(secoes_2026: pl.DataFrame) -> None:
    assert lb.locais_do_bairro(secoes_2026, "LAP").is_empty()
    assert lb.locais_do_bairro(secoes_2026, "LAP", contem=True)["NR_LOCAL_VOTACAO"].to_list() == [1040]
    assert lb.sugestoes(secoes_2026, "lap") == ["LAPA"]


def test_fonte_ibge_usa_a_coordenada_e_nao_o_texto(tse_cache: Path, secoes_2026: pl.DataFrame) -> None:
    escrever_bairros(tse_cache)
    geo = br.malha("RJ", tse_cache)
    # a Escola X de 2026 diz "CENTRO" no cadastro, mas a coordenada cai na Lapa Sintética da malha;
    # a Creche Nova não tem coordenada e fica de fora
    lapa = lb.locais_do_bairro(secoes_2026, "Lapa Sintetica", fonte="ibge", geo=geo)
    assert lapa.select("NR_LOCAL_VOTACAO", "BAIRRO").rows() == [(1023, "Lapa Sintética")]
    assert lb.locais_do_bairro(secoes_2026, "Centro Sintético", fonte="ibge", geo=geo)["NR_LOCAL_VOTACAO"].to_list() == [1015]


def test_fonte_ibge_sem_malha_falha(secoes_2026: pl.DataFrame) -> None:
    with pytest.raises(ValueError, match="malha"):
        lb.locais_do_bairro(secoes_2026, "Centro", fonte="ibge")


def test_tabela_locais_rotula_agregada(secoes_2026: pl.DataFrame) -> None:
    locais = lb.locais_do_bairro(secoes_2026, "centro")
    tab = lb.tabela_locais(locais, lb.secoes_dos_locais(secoes_2026, locais))
    assert not [c for c in tab.columns if c.startswith("_")]
    pedro_rio = tab.filter((pl.col("NM_MUNICIPIO") == "RIO DE JANEIRO") & (pl.col("NR_LOCAL_VOTACAO") == 1015))
    assert pedro_rio.select("N_SECOES", "QT_ELEITORES", "SECOES").row(0) == (4, 955, "10, 11, 12, 13 (agregada à 12)")


def test_cli_imprime_e_grava_planilha(tse_cache: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    saida = tmp_path / "out" / "centro.xlsx"
    rc = cli.main(["CENTRO", "--ano", "2026", "--municipio", "Rio de Janeiro", "--cache-dir", str(tse_cache),
                   "--saida", str(saida)])
    out = capsys.readouterr().out
    assert rc == 0
    assert f"Cadastro de eleitorado 2026, 1º turno — versão {CARIMBO}" in out
    assert "Zona 4 · Local 1015 — COLEGIO PEDRO II" in out
    assert "Seções: 10, 11, 12, 13 (agregada à 12)" in out
    assert "4 seções, 955 eleitores" in out
    assert "Total: 2 locais, 5 seções, 1.105 eleitores" in out
    assert "NITERÓI" not in out
    assert f"Planilha gravada em {saida}" in out
    wb = openpyxl.load_workbook(saida, read_only=True)
    assert wb.sheetnames == ["Locais", "Seções", "Consulta"]
    aba = {nome: list(wb[nome].iter_rows(values_only=True)) for nome in wb.sheetnames}
    col = aba["Locais"][0].index("NR_LOCAL_VOTACAO")
    assert [r[col] for r in aba["Locais"][1:]] == [1015, 1023]
    tipo = aba["Seções"][0].index("TIPO_SECAO")
    assert [r[tipo] for r in aba["Seções"][1:]].count("AGREGADA") == 1 and len(aba["Seções"]) == 6
    consulta = dict(aba["Consulta"][1:])
    assert consulta["bairro"] == "CENTRO" and consulta["eleitores"] == "1105" and consulta["turno"] == "1"
    # a versão do cadastro usada vai para a Consulta e fica guardada no acervo
    prov = json.loads((tse_cache / ZIP_2026).with_suffix(".proveniencia.json").read_text())
    assert consulta["versão do arquivo (carimbo)"] == CARIMBO and consulta["SHA-512"] == prov["sha512"]
    assert consulta["geração do TSE"] == "29/09/2026 06:28:09"
    assert [x.carimbo for x in vs.listar(tse_cache, ZIP_2026)] == [CARIMBO]
    assert not list(saida.parent.glob("*.tmp"))


def test_destino_padrao_leva_turno_fonte_e_versao(tse_cache: Path) -> None:
    el.ensure_local_provenance(v.electorate_spec(2026, "RJ"), tse_cache)
    versao = vs.em_uso(tse_cache, ZIP_2026)
    assert cli.destino_padrao("RJ", "Barra da Tijuca", 2026, 1, versao) == \
        Path(f"saidas/RJ/locais_bairro_barradatijuca_2026_t1_{CARIMBO}.xlsx")
    assert cli.destino_padrao("RJ", "Barra da Tijuca", 2026, 2, versao, "ibge") == \
        Path(f"saidas/RJ/locais_bairro_barradatijuca_2026_t2_ibge_{CARIMBO}.xlsx")


def test_cli_grava_no_destino_padrao(tse_cache: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    assert cli.main(["CENTRO", "--ano", "2026", "--municipio", "Rio de Janeiro",
                     "--cache-dir", str(tse_cache)]) == 0
    assert (tmp_path / f"saidas/RJ/locais_bairro_centro_2026_t1_{CARIMBO}.xlsx").is_file()


def test_versao_guardada_refaz_a_lista_sem_o_zip_no_cache(tse_cache: Path, tmp_path: Path) -> None:
    primeira, refeita = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    args = ["CENTRO", "--ano", "2026", "--municipio", "Rio de Janeiro", "--cache-dir", str(tse_cache)]
    assert cli.main([*args, "--saida", str(primeira)]) == 0
    for f in tse_cache.glob("eleitorado_local_votacao_2026*"):  # o cache perde o ZIP; o acervo continua com ele
        f.unlink()
    assert cli.main([*args, "--versao", "2026-09", "--saida", str(refeita)]) == 0
    abas = [{n: list(openpyxl.load_workbook(x, read_only=True)[n].iter_rows(values_only=True))
             for n in ("Locais", "Seções")} for x in (primeira, refeita)]
    assert abas[0] == abas[1]


def test_versao_inexistente_lista_as_guardadas(tse_cache: Path, caplog: pytest.LogCaptureFixture) -> None:
    assert cli.main(["CENTRO", "--ano", "2026", "--cache-dir", str(tse_cache), "--versao", "1999"]) == 1
    assert "nenhuma versão guardada" in caplog.text


def test_turno_sem_linhas_no_cadastro_e_erro(tse_cache: Path, tmp_path: Path,
                                            caplog: pytest.LogCaptureFixture) -> None:
    saida = tmp_path / "t2.xlsx"
    # o cadastro sintético só tem o 1º turno
    assert cli.main(["CENTRO", "--ano", "2026", "--turno", "2", "--cache-dir", str(tse_cache),
                     "--saida", str(saida)]) == 1
    assert "não tem linhas do 2º turno" in caplog.text and not saida.exists()


def test_cli_sem_resultado_sugere_e_nao_grava(tse_cache: Path, tmp_path: Path, caplog: pytest.LogCaptureFixture) -> None:
    saida = tmp_path / "lap.xlsx"
    assert cli.main(["Lap", "--ano", "2026", "--cache-dir", str(tse_cache), "--saida", str(saida)]) == 1
    assert "LAPA" in caplog.text
    assert not saida.exists()
