"""Áreas de ponderação do Censo 2022 no Perfil × voto (apuracao.areas_ponderacao, rodada 49) — offline."""

from __future__ import annotations

import io
import json
from pathlib import Path

import polars as pl
import pytest
from fastapi.testclient import TestClient

from apuracao import areas_ponderacao as ap
from apuracao import bairros as br
from apuracao import perfil as pf
from apuracao import perfil_local as pfl
from apuracao.web.app import create_app
from conftest import (AREA_NIT, AREA_RIO_1, AREA_RIO_2, CAND, download_sem_rede, escrever_areas, escrever_bairros,
                      escrever_setores)

PEDRO, ESCOLA_X, CIEP, NITEROI = "3304557000401015", "3304557000401023", "3304557000401031", "3303302007101015"


@pytest.fixture(autouse=True)
def _sem_rede(monkeypatch: pytest.MonkeyPatch) -> None:
    download_sem_rede(monkeypatch)


def _xlsx(linhas: list[list]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    for r in linhas:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


RELIGIAO = [["Censo Demográfico 2022 - Resultados Gerais da Amostra por Áreas de Ponderação"],
            ["Tabela 4_1 - Pessoas de 10 anos ou mais de idade, por religião"],
            ["Código do Município", "Nome do Município", "Área de ponderação", "Nome", "Total",
             "Católica Apostólica Romana", "Evangélicas", "Espírita", "Umbanda e Candomblé", "Tradições indígenas",
             "Outras religiosidades", "Sem religião", "Não sabe", "Sem declaração"],
            ["Brasil", "Brasil", "Brasil", "Brasil", 300, 150, 90, 10, 2, 0, 8, 40, 0, 0],
            [3304557, "Rio de Janeiro", 3304557001, "Área 001", 200, 100, 60, 10, 2, "-", 8, 20, "-", "-"],
            [3303302, "Niterói", 3303302001, "Área 001", 100, 50, 30, "-", "X", "-", "-", 20, "-", "-"]]


# --------------------------------------------------------------------------- leitura das tabelas do IBGE
def test_ler_tabela_da_amostra() -> None:
    df, brasil = ap.ler_tabela(_xlsx(RELIGIAO), "Tab4_1.xlsx")
    rj = df.filter(pl.col("CD_AP") == "3304557001").row(0, named=True)
    assert (rj["CD_MUN"], rj["NM_AP"], rj["c4"], rj["c6"], rj["c11"]) == (3304557, "Área 001", 200, 60, 20)
    nit = df.filter(pl.col("CD_AP") == "3303302001").row(0, named=True)
    assert nit["c7"] == 0 and nit["c8"] is None  # "-" = zero; "X" (sigilo) = sem dado
    assert brasil[4] == 300 and brasil[6] == 90
    conf = ap.conferir_brasil({"Tab4_1.xlsx": df}, {"Tab4_1.xlsx": brasil})
    assert conf.filter(pl.col("COLUNA") == "Total")["DIF_PCT"].item() == 0


def test_cabecalho_mudado_no_ibge_e_erro() -> None:
    mudado = [r[:] for r in RELIGIAO]
    mudado[2][6] = "Evangélica"  # o IBGE renomeou a coluna: nada de número na coluna errada
    with pytest.raises(ValueError, match="cabeçalho mudou"):
        ap.ler_tabela(_xlsx(mudado), "Tab4_1.xlsx")


def test_indicadores_da_amostra() -> None:
    rel, _ = ap.ler_tabela(_xlsx(RELIGIAO), "Tab4_1.xlsx")
    tabelas = {n: rel if n == "Tab4_1.xlsx" else rel.select("CD_AP", "CD_MUN", "NM_MUN", "NM_AP").with_columns(
        *[pl.lit(10.0).alias(f"c{i}") for i in ap.CABECALHOS[n]]) for n in ap.CABECALHOS}
    ind = {r["CD_AP"]: r for r in ap.indicadores_amostra(tabelas).iter_rows(named=True)}
    rj = ind["3304557001"]
    assert (rj["pct_catolicos"], rj["pct_evangelicos"], rj["pct_sem_religiao"]) == (50, 30, 10)
    assert ind["3303302001"]["pct_matriz_africana"] is None  # sigilo no numerador
    assert rj["renda_pc_media"] == 10  # valor (R$), não percentual
    assert rj["pct_trabalho_1h_mais"] == 300  # três colunas de 10 sobre o total de 10 (dado sintético)


def test_ler_composicao() -> None:
    comp = ap.ler_composicao(_xlsx([["Setor", "Código do Município", "Área de Ponderação", "Nome do Municíipio"],
                                    [330455705000001, 3304557, 3304557001, "Rio de Janeiro"]]))
    assert comp.row(0) == ("330455705000001", "3304557001", 3304557)
    with pytest.raises(ValueError):
        ap.ler_composicao(_xlsx([["Setor", "Município", "Área"]]))


# --------------------------------------------------------------------------- perfil × voto por área
@pytest.fixture()
def cache(tse_cache: Path) -> Path:
    escrever_bairros(tse_cache)
    escrever_setores(tse_cache)
    escrever_areas(tse_cache)
    return tse_cache


@pytest.fixture()
def pva(cache: Path, monkeypatch: pytest.MonkeyPatch) -> ap.PerfilVotoArea:
    comp = br.ComparacaoBairros(br.Bairros("RJ", cache))
    monkeypatch.setattr(comp, "siglas", lambda ano: {55: "PSD"})
    return ap.PerfilVotoArea(comp)


def test_votos_e_perfil_por_area(pva: ap.PerfilVotoArea, cache: Path) -> None:
    assert dict(pva.local_area(2024).iter_rows()) == {PEDRO: AREA_RIO_1, ESCOLA_X: AREA_RIO_2, CIEP: AREA_RIO_2,
                                                      NITEROI: AREA_NIT}
    y = pf.Alvo(2024, 13, numero=CAND)
    local = pfl.PerfilVotoLocal(pva.comp).voto(y)
    soma = lambda us: local.filter(pl.col("CD_BAIRRO").is_in(us)).select("VOTOS", "VALIDOS").sum().row(0)  # noqa: E731
    por_area = {r["CD_BAIRRO"]: r for r in pva.voto(y).iter_rows(named=True)}
    assert (por_area[AREA_RIO_2]["VOTOS"], por_area[AREA_RIO_2]["VALIDOS"]) == soma([ESCOLA_X, CIEP])
    tse_local = {r["CD_BAIRRO"]: r for r in pfl.PerfilVotoLocal(pva.comp).perfil(2024).iter_rows(named=True)}
    p = {r["CD_BAIRRO"]: r for r in pva.perfil(2024).iter_rows(named=True)}
    e, c = tse_local[ESCOLA_X], tse_local[CIEP]
    assert p[AREA_RIO_2]["ELEITORES"] == e["ELEITORES"] + c["ELEITORES"]
    assert p[AREA_RIO_2]["pct_superior"] == pytest.approx(
        (e["pct_superior"] * e["ELEITORES"] + c["pct_superior"] * c["ELEITORES"]) / p[AREA_RIO_2]["ELEITORES"])


def test_censo_por_area(pva: ap.PerfilVotoArea) -> None:
    c = {r["CD_BAIRRO"]: r for r in pva.censo_local(2024).iter_rows(named=True)}
    assert c[AREA_RIO_1]["pct_evangelicos"] == 10 and c[AREA_RIO_2]["pct_evangelicos"] == 40  # amostra
    assert c[AREA_RIO_2]["pct_apartamentos"] == pytest.approx((50 + 20) / 2)  # universo: setores da área
    assert c[AREA_RIO_1]["renda_media"] == 8000  # e os indicadores do local, agora por área
    ind = pva.indicadores()
    assert ind["pct_evangelicos"]["fonte"] == ap.FONTE_AMOSTRA and "pct_evangelicos" not in pfl.PerfilVotoLocal.indicadores()


def test_dispersao_por_area(pva: ap.PerfilVotoArea) -> None:
    d = pva.dispersao(pf.Alvo(2024, 13, numero=CAND), "pct_evangelicos", min_validos=0)
    pts = {r["CD_BAIRRO"]: r for r in d["pontos"].iter_rows(named=True)}
    assert set(pts) <= {AREA_RIO_1, AREA_RIO_2, AREA_NIT} and AREA_RIO_2 in pts
    assert pts[AREA_RIO_2]["BAIRRO"] == "Área 002 — Rio de Janeiro"
    assert {r["CD_MUN"] for r in pva.municipios().iter_rows(named=True)} == {3304557, 3303302}


# --------------------------------------------------------------------------- API
def test_api_por_area(cache: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    (cache / "malhas" / "municipios_RJ.geojson").write_text(json.dumps({"type": "FeatureCollection", "features": []}))
    monkeypatch.setattr(br.ComparacaoBairros, "siglas", lambda self, ano: {55: "PSD"})
    site = TestClient(create_app(tmp_path / "dados", "RJ", cache))
    info = site.get("/api/perfil/info?unidade=area").json()
    assert info["unidade"] == "area" and info["indicadores"]["pct_evangelicos"]["fonte"] == ap.FONTE_AMOSTRA
    base = f"ano=2024&cargo=13&numero={CAND}&min_validos=0&unidade=area"
    d = site.get(f"/api/perfil/dispersao?{base}&x=pct_evangelicos").json()
    assert d["estatistica"]["n"] == len(d["pontos"]) >= 2
    assert len(site.get(f"/api/perfil/correlacoes?{base}").json()["correlacoes"]) == len(ap.PerfilVotoArea.indicadores())


# --------------------------------------------------------------------------- mapa por área (TODO 25)
def test_malha_das_areas_funde_os_setores(cache: Path) -> None:
    geo = ap.malha("RJ", cache)
    props = {f["properties"]["CD_AP"]: f["properties"] for f in geo["features"]}
    assert set(props) == {AREA_RIO_1, AREA_RIO_2, AREA_NIT}  # 5 setores → 3 áreas
    assert props[AREA_RIO_2]["NM_AP"] == "Área 002" and props[AREA_NIT]["CD_MUN"] == 3303302
    assert (cache / ap.MALHA.format(uf="RJ")).exists() and ap.malha("RJ", cache) == geo  # 2ª vez: do cache


def test_valores_do_mapa_por_area(pva: ap.PerfilVotoArea) -> None:
    y = pf.Alvo(2024, 13, numero=CAND)
    voto = {r["CD_BAIRRO"]: r["VOTO"] for r in pva.voto(y).iter_rows(named=True)}
    d = ap.mapa(pva, 2024, "voto", cargo=13, metrica="pct_candidato", numero=CAND)
    assert d["tipo"] == "sequencial" and d["unidade"] == "%" and d["metrica"] == "pct_candidato"
    assert {a: i["valor"] for a, i in d["itens"].items()} == pytest.approx(voto)
    assert d["itens"][AREA_RIO_2]["municipio"] == "Área 002 — Rio de Janeiro" and d["cobertura"]["areas"] == 3
    p = ap.mapa(pva, 2024, "perfil", indicador="pct_evangelicos")
    assert p["itens"][AREA_RIO_2]["valor"] == 40 and p["fonte_indicador"] == ap.FONTE_AMOSTRA and p["metrica"] is None
    rio = ap.mapa(pva, 2024, "perfil", indicador="pct_evangelicos", municipio=3304557)
    assert set(rio["itens"]) == {AREA_RIO_1, AREA_RIO_2} and rio["cobertura"]["areas"] == 2
    with pytest.raises(ValueError):
        ap.mapa(pva, 2024, "setor")  # camada desconhecida
    with pytest.raises(ValueError):
        ap.mapa(pva, 2024, "residuo", indicador="pct_evangelicos")  # o resíduo precisa do número


def test_abstencao_por_area_soma_os_aptos(pva: ap.PerfilVotoArea) -> None:
    from apuracao import mapa_locais as ml
    loc = ml._participacao(pva, 2024, 13, 1, "abstencao_pct")
    e = loc.filter(pl.col("CD_BAIRRO").is_in([ESCOLA_X, CIEP])).select("NUM", "DEN").sum().row(0)
    d = ap.mapa(pva, 2024, "voto", cargo=13, metrica="abstencao_pct")
    assert d["itens"][AREA_RIO_2]["valor"] == pytest.approx(100 * e[0] / e[1])  # Σ abstenções ÷ Σ aptos


def test_api_do_mapa_por_area(cache: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    (cache / "malhas" / "municipios_RJ.geojson").write_text(json.dumps({"type": "FeatureCollection", "features": []}))
    monkeypatch.setattr(br.ComparacaoBairros, "siglas", lambda self, ano: {55: "PSD"})
    site = TestClient(create_app(tmp_path / "dados", "RJ", cache))
    assert len(site.get("/geo/areas.geojson").json()["features"]) == 3
    d = site.get(f"/api/mapa/areas?ano=2024&camada=voto&cargo=13&metrica=pct_candidato&numero={CAND}").json()
    assert set(d["itens"]) <= {AREA_RIO_1, AREA_RIO_2, AREA_NIT} and d["itens"]
    assert site.get("/api/mapa/areas?ano=2024&camada=perfil&indicador=pct_evangelicos").json()["itens"][AREA_NIT]["valor"] == 25
    assert site.get("/api/mapa/areas?ano=2024&camada=perfil&indicador=xyz").status_code == 400
    r = site.post("/api/exportar/mapa", json={"formato": "svg", "camada": "areas", "titulo": "t",
                                              "cores": {AREA_RIO_1: "#ff0000"}, "legenda": [["#ff0000", "x"]]})
    assert r.status_code == 200 and b"<svg" in r.content[:400]


def test_residuo_por_area_divergente(pva: ap.PerfilVotoArea, monkeypatch: pytest.MonkeyPatch) -> None:
    # 3 áreas sintéticas não dão reta (mínimo de 5 unidades): sem resíduo, mapa vazio e a estatística diz por quê
    vazio = ap.mapa(pva, 2024, "residuo", cargo=13, numero=CAND, indicador="pct_evangelicos", min_validos=0)
    assert vazio["itens"] == {} and vazio["estatistica"]["n"] == 3 and vazio["tipo"] == "divergente"
    pontos = pl.DataFrame({"CD_BAIRRO": [AREA_RIO_1, AREA_RIO_2], "RESIDUO": [-4.0, 6.5], "Y": [30.0, 52.0],
                           "X": [10.0, 40.0]})
    monkeypatch.setattr(pva, "dispersao", lambda y, x, mv, pond, mun: {
        "pontos": pontos, "estatistica": {"n": 2, "pearson": 0.9, "r2": 0.81, "a": 1.0, "b": 2.0, "p": 0.1},
        "rotulo_x": "% de evangélicos", "rotulo_y": "% dos válidos"})
    d = ap.mapa(pva, 2024, "residuo", cargo=13, numero=CAND, indicador="pct_evangelicos", min_validos=0)
    assert d["unidade"] == "p.p." and d["sentido"] == "residuo" and d["estatistica"]["pearson"] == 0.9
    assert {a: i["valor"] for a, i in d["itens"].items()} == {AREA_RIO_1: -4.0, AREA_RIO_2: 6.5}
    assert (d["itens"][AREA_RIO_2]["voto"], d["itens"][AREA_RIO_2]["indicador"]) == (52.0, 40.0)  # para a dica


def test_variacao_por_area(pva: ap.PerfilVotoArea, monkeypatch: pytest.MonkeyPatch) -> None:
    """A área é a mesma nos dois anos: a variação é a métrica de 2024 − a de 2020, área por área (p.p.)."""
    abst = {2020: {AREA_RIO_1: 20.0, AREA_RIO_2: 30.0}, 2024: {AREA_RIO_1: 25.0, AREA_RIO_2: 22.5, AREA_NIT: 10.0}}
    monkeypatch.setattr(pva, "participacao", lambda ano, cargo, turno, metrica: pl.DataFrame(
        {"CD_BAIRRO": list(abst[ano]), "VALOR": list(abst[ano].values())}).with_columns(
        pl.lit(1).alias("NUM"), pl.lit(1).alias("DEN")))
    d = ap.mapa(pva, 2024, "variacao", cargo=13, metrica="abstencao_pct", ano_ref=2020)
    assert d["tipo"] == "divergente" and d["lados"] == ["2020", "2024"] and d["sentido"] == "variacao"
    assert {a: i["valor"] for a, i in d["itens"].items()} == {AREA_RIO_1: 5.0, AREA_RIO_2: -7.5}  # Niterói: sem 2020
    assert (d["itens"][AREA_RIO_2]["antes"], d["itens"][AREA_RIO_2]["depois"]) == (30.0, 22.5)
    with pytest.raises(ValueError):
        ap.mapa(pva, 2024, "variacao", cargo=13, metrica="abstencao_pct", ano_ref=2024)


def test_destino_dos_eliminados_por_area(pva: ap.PerfilVotoArea, monkeypatch: pytest.MonkeyPatch) -> None:
    """A inferência por local somada na área: % observados e previstos ponderados pelos aptos do turno; o destino
    dos eliminados (estimado por município) é o do município."""
    from types import SimpleNamespace

    from apuracao import mapa_locais as ml
    a = "A (X)"
    por = pl.DataFrame({
        "UNIDADE": [PEDRO, ESCOLA_X, CIEP, NITEROI], "APTOS_1": [100, 300, 100, 50], "APTOS_2": [100, 300, 100, 50],
        "ABST_1_PCT": [10.0, 20.0, 40.0, 30.0], "ABST_2_PCT": [15.0, 22.0, 50.0, 30.0],
        "ELIMINADOS_1_PCT": [30.0, 10.0, 50.0, 20.0], f"{a}_2_PCT": [60.0, 40.0, 20.0, 50.0],
        f"{a}_2_AJUSTE_PCT": [55.0, 42.0, 30.0, 50.0], f"ELIM_PARA_{a}_PCT": [70.0, 70.0, 70.0, 35.0],
        "ELIM_PARA_B (Y)_PCT": [30.0, 30.0, 30.0, 65.0]})
    res = SimpleNamespace(unidades=SimpleNamespace(cat2=[a, "B (Y)"]))
    monkeypatch.setattr(ml, "transferencia_por_local", lambda unidade, ano, cargo: (res, por))
    _, area = ap.transferencia_por_area(pva, 2024, 7)
    r2 = area.filter(pl.col("UNIDADE") == AREA_RIO_2).row(0, named=True)  # Escola X (300 aptos) + CIEP (100)
    assert r2["ABST_1_PCT"] == pytest.approx((20 * 300 + 40 * 100) / 400)
    assert r2["ABST_EXTRA_PP"] == pytest.approx((22 * 300 + 50 * 100) / 400 - (20 * 300 + 40 * 100) / 400)
    assert r2["RESIDUO_A_PP"] == pytest.approx((40 * 300 + 20 * 100) / 400 - (42 * 300 + 30 * 100) / 400)
    assert r2[f"ELIM_PARA_{a}_PCT"] == 70  # o do município (Rio)
    d = ap.mapa(pva, 2024, "transferencia", cargo=7, metrica="elim_para_a")
    assert d["tipo"] == "sequencial" and d["itens"][AREA_NIT]["valor"] == 35 and d["finalistas"] == [a, "B (Y)"]
    e = ap.mapa(pva, 2024, "transferencia", cargo=7, metrica="abst_extra", municipio=3304557)
    assert e["tipo"] == "divergente" and set(e["itens"]) == {AREA_RIO_1, AREA_RIO_2} and e["lados"] == ["1º turno", "2º turno"]
    assert e["itens"][AREA_RIO_1]["valor"] == pytest.approx(5) and e["itens"][AREA_RIO_1]["antes"] == 10
