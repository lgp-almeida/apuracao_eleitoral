"""Diferenças entre duas versões do cadastro de eleitorado (`apuracao/diferencas_cadastro.py` e
`ingerir_eleitorado.py --diferencas`) — offline.

Cenário sintético (Rio, zona 4; Niterói, zona 71), versão de antes gerada em 16/10/2026 06:29:01 e a de depois em
17/10/2026 06:28:28, as duas com os dois turnos:

  Rio 10  1015 nos dois turnos, nas duas versões                        → sem diferença
  Rio 11  1º turno no 1050 (temporário; original 1015); no 2º, antes
          voltava ao 1015, depois fica no 1050                          → VOLTA_AO_ORIGINAL → —; 5 células
  Rio 12  1023; depois, o endereço do 2º turno ganha "(ACESSO …)"       → — → NOME_OU_ENDERECO; 1 célula
  Rio 14  só antes                                                      → extinta (2 linhas)
  Rio 30  só depois                                                     → criada (2 linhas)
  Nit 11  depois, a latitude do local muda nos dois turnos              → 2 células
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import polars as pl
import pytest

import ingerir_eleitorado as cli
import votos_por_local_votacao as v
from apuracao import diferencas_cadastro as dc
from apuracao import eleitorado as el
from apuracao import versoes as vs
from conftest import EL_HEADER, NIT, RIO, SAO, _csv

SPEC = v.electorate_spec(2026, "RJ")
LOCAIS = {  # nº → (nome, endereço, bairro, latitude, longitude)
    1015: ("COLÉGIO PEDRO II", "AV. MARECHAL FLORIANO, 80", "CENTRO", "-22,9035", "-43,1790"),
    1023: ("ESCOLA X", "RUA DO LAVRADIO, 12", "LAPA", "-22,9130", "-43,1800"),
    1050: ("GINÁSIO PROVISÓRIO", "RUA W, 9", "CENTRO", "-22,9100", "-43,1700"),
}
ACESSO = "RUA DO LAVRADIO, 12 (ACESSO PELA RUA X)"
# (turno, município, zona, seção, local, local original, eleitores, outro endereço, outra latitude)
ANTES = [
    (1, RIO, 4, 10, 1015, 1015, 300, None, None), (2, RIO, 4, 10, 1015, 1015, 300, None, None),
    (1, RIO, 4, 11, 1050, 1015, 280, None, None), (2, RIO, 4, 11, 1015, 1015, 280, None, None),
    (1, RIO, 4, 12, 1023, 1023, 250, None, None), (2, RIO, 4, 12, 1023, 1023, 250, None, None),
    (1, RIO, 4, 14, 1023, 1023, 100, None, None), (2, RIO, 4, 14, 1023, 1023, 100, None, None),
    (1, NIT, 71, 11, 1015, 1015, 310, None, None), (2, NIT, 71, 11, 1015, 1015, 310, None, None),
]
DEPOIS = [
    (1, RIO, 4, 10, 1015, 1015, 300, None, None), (2, RIO, 4, 10, 1015, 1015, 300, None, None),
    (1, RIO, 4, 11, 1050, 1015, 280, None, None), (2, RIO, 4, 11, 1050, 1015, 280, None, None),
    (1, RIO, 4, 12, 1023, 1023, 250, None, None), (2, RIO, 4, 12, 1023, 1023, 250, ACESSO, None),
    (1, RIO, 4, 30, 1015, 1015, 90, None, None), (2, RIO, 4, 30, 1015, 1015, 90, None, None),
    (1, NIT, 71, 11, 1015, 1015, 310, None, "-22,8950"), (2, NIT, 71, 11, 1015, 1015, 310, None, "-22,8950"),
]
GERACAO = {"antes": ("16/10/2026", "06:29:01"), "depois": ("17/10/2026", "06:28:28")}
CARIMBO = {"antes": "2026-10-16T062901", "depois": "2026-10-17T062828"}


def _linha(geracao: tuple[str, str], dado: tuple, uf: str = "RJ") -> list:
    turno, (cd, nm), zona, secao, local, original, eleitores, endereco, lat = dado
    nome_local, end_local, bairro, lat_local, lon = LOCAIS[local]
    campos = dict.fromkeys(EL_HEADER, "")
    campos.update(
        DT_GERACAO=geracao[0], HH_GERACAO=geracao[1], AA_ELEICAO=2026, NR_TURNO=turno, SG_UF=uf,
        DT_ELEICAO="04/10/2026" if turno == 1 else "25/10/2026", DS_ELEICAO=f"{turno}º Turno",
        CD_MUNICIPIO=str(cd), NM_MUNICIPIO=nm, NR_ZONA=zona, NR_SECAO=secao, NR_SECAO_PRINCIPAL=-1,
        CD_TIPO_SECAO_AGREGADA=1, DS_TIPO_SECAO_AGREGADA="Principal", NR_LOCAL_VOTACAO=local,
        NM_LOCAL_VOTACAO=nome_local, DS_ENDERECO=endereco or end_local, NM_BAIRRO=bairro,
        NR_LATITUDE=lat or lat_local, NR_LONGITUDE=lon, DS_SITU_LOCAL_VOTACAO="ATIVO", QT_ELEITOR_SECAO=eleitores,
        NR_LOCAL_VOTACAO_ORIGINAL=original, NM_LOCAL_VOTACAO_ORIGINAL=LOCAIS[original][0],
        DS_ENDERECO_LOCVT_ORIGINAL=LOCAIS[original][1],
    )
    return [campos[c] for c in EL_HEADER]


def escrever(cache: Path, versao: str) -> Path:
    """O ZIP da versão no cache, trocado como o download troca (o do acervo é só leitura)."""
    cache.mkdir(parents=True, exist_ok=True)
    linhas = [_linha(GERACAO[versao], d) for d in (ANTES if versao == "antes" else DEPOIS)]
    tmp = cache / f"{SPEC.zip_name}.tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("eleitorado_local_votacao_2026_RJ.csv", _csv(EL_HEADER, linhas))
        zf.writestr("eleitorado_local_votacao_2026_BRASIL.csv", _csv(EL_HEADER, linhas))
        zf.writestr("leiame.pdf", b"%PDF-fake")
    tmp.replace(cache / SPEC.zip_name)
    prov = (cache / SPEC.zip_name).with_suffix(".proveniencia.json")
    prov.unlink(missing_ok=True)  # o ZIP é outro: a proveniência é registrada de novo
    for p in cache.glob("eleitorado_local_votacao_2026__*.parquet"):
        p.unlink()
    return cache


@pytest.fixture
def acervo(tmp_path: Path) -> Path:
    """Cache com a versão de antes guardada no acervo e a de depois no cache (ainda não guardada)."""
    cache = escrever(tmp_path / "cache", "antes")
    el.ensure_local_provenance(SPEC, cache)
    vs.em_uso(cache, SPEC.zip_name)
    return escrever(cache, "depois")


@pytest.fixture
def versoes(acervo: Path) -> tuple[vs.Versao, vs.Versao]:
    el.ensure_local_provenance(SPEC, acervo)
    return vs.escolher(acervo, SPEC.zip_name, CARIMBO["antes"]), vs.em_uso(acervo, SPEC.zip_name)


@pytest.fixture
def diferencas(versoes: tuple[vs.Versao, vs.Versao]) -> dc.Diferencas:
    return dc.comparar_zips(versoes[0].zip, versoes[1].zip, ["RJ"])


def _turnos(versao: vs.Versao) -> dc.Turnos:
    return dc.turnos(*[el.load_sections(2026, "RJ", versao.pasta, turno=t) for t in (1, 2)])


def test_celulas_criadas_e_extintas(diferencas: dc.Diferencas) -> None:
    d = diferencas
    assert (d.linhas_antes, d.linhas_depois) == (10, 10)
    assert d.criadas.select("NR_TURNO", "NR_SECAO", "NR_LOCAL_VOTACAO").rows() == [(1, 30, "1015"), (2, 30, "1015")]
    assert d.extintas.select("NR_TURNO", "NR_SECAO").rows() == [(1, 14), (2, 14)]
    assert d.colunas_novas == d.colunas_removidas == ()
    rio_11 = d.celulas.filter((pl.col("CD_MUNICIPIO") == RIO[0]) & (pl.col("NR_SECAO") == 11))
    assert set(rio_11["NR_TURNO"]) == {2}
    assert sorted(rio_11["COLUNA"]) == ["DS_ENDERECO", "NM_LOCAL_VOTACAO", "NR_LATITUDE", "NR_LOCAL_VOTACAO",
                                        "NR_LONGITUDE"]  # o bairro é o mesmo (CENTRO) e o original não muda
    troca = rio_11.filter(pl.col("COLUNA") == "NR_LOCAL_VOTACAO").row(0, named=True)
    assert (troca["ANTES"], troca["DEPOIS"], troca["NR_LOCAL_VOTACAO"], troca["NM_MUNICIPIO"]) == \
        ("1015", "1050", 1050, "RIO DE JANEIRO")
    assert d.celulas.height == 8 and d.celulas.select(dc.CHAVE).n_unique() == 4
    assert not {"DT_GERACAO", "HH_GERACAO"} & set(d.celulas["COLUNA"])  # a hora da geração não conta


def test_resumo_por_coluna_e_alteracoes_por_local(diferencas: dc.Diferencas) -> None:
    r = {x["COLUNA"]: (x["LINHAS"], x["LINHAS_1T"], x["LINHAS_2T"])
         for x in dc.resumo_colunas(diferencas).iter_rows(named=True)}
    assert r == {"NR_LATITUDE": (3, 1, 2), "DS_ENDERECO": (2, 0, 2), "NM_LOCAL_VOTACAO": (1, 0, 1),
                 "NR_LOCAL_VOTACAO": (1, 0, 1), "NR_LONGITUDE": (1, 0, 1)}
    assert list(r)[:2] == ["NR_LATITUDE", "DS_ENDERECO"]  # a mais alterada primeiro
    alt = dc.alteracoes(diferencas)
    acesso = alt.filter((pl.col("NR_LOCAL_VOTACAO") == 1023) & (pl.col("COLUNA") == "DS_ENDERECO")).row(0, named=True)
    assert (acesso["NR_TURNO"], acesso["ANTES"], acesso["DEPOIS"], acesso["SECOES"]) == \
        (2, "RUA DO LAVRADIO, 12", ACESSO, "12")
    assert alt.filter(pl.col("NM_MUNICIPIO") == "NITERÓI")["NR_TURNO"].to_list() == [1, 2]


def test_situacoes_e_mudancas_entre_as_versoes(versoes: tuple[vs.Versao, vs.Versao], diferencas: dc.Diferencas) -> None:
    antes, depois = (_turnos(x) for x in versoes)
    lado = {x["SITUACAO"]: x for x in dc.situacoes(antes, depois).iter_rows(named=True)}
    assert (lado["VOLTA_AO_ORIGINAL"]["SECOES_ANTES"], lado["VOLTA_AO_ORIGINAL"]["SECOES_DEPOIS"]) == (1, 0)
    assert (lado["NOME_OU_ENDERECO"]["SECOES_ANTES"], lado["NOME_OU_ENDERECO"]["ELEITORES_DEPOIS"]) == (0, 250)
    m = dc.mudancas_de_situacao(antes, depois)
    assert m is not None  # as criadas e as extintas (30 e 14) não entram: estão nas diferenças do CSV
    assert m.select("NR_SECAO", "SITUACAO_ANTES", "SITUACAO_DEPOIS", "NR_LOCAL_1T", "NR_LOCAL_2T_ANTES",
                    "NR_LOCAL_2T_DEPOIS").rows() == [
        (11, "VOLTA_AO_ORIGINAL", dc.SEM_SITUACAO, 1050, 1015, 1050),
        (12, dc.SEM_SITUACAO, "NOME_OU_ENDERECO", 1023, 1023, 1023)]
    assert dc.quadro_de_situacoes(m).select("SITUACAO_ANTES", "SECOES", "ELEITORES", "UFS").rows() == [
        ("VOLTA_AO_ORIGINAL", 1, 280, "RJ"), (dc.SEM_SITUACAO, 1, 250, "RJ")]
    uf = dc.resumo_ufs(diferencas, antes, depois).row(0, named=True)
    assert (uf["LINHAS_ALTERADAS"], uf["CELULAS"], uf["CRIADAS"], uf["EXTINTAS"]) == (4, 8, 2, 2)
    assert (uf["TEMPORARIAS_1T_ANTES"], uf["TEMPORARIAS_2T_ANTES"], uf["TEMPORARIAS_2T_DEPOIS"]) == (1, 0, 1)


def test_versao_sem_2o_turno_nao_tem_situacao(versoes: tuple[vs.Versao, vs.Versao]) -> None:
    so_1t = dc.turnos(*[el.load_sections(2026, "RJ", versoes[0].pasta, turno=1)] * 2).t1.head(0)
    sem = dc.Turnos(so_1t, so_1t, None)
    assert dc.mudancas_de_situacao(sem, _turnos(versoes[1])) is None
    assert dc.situacoes(sem, _turnos(versoes[1]))["SECOES_ANTES"].null_count() == 8


def test_csv_nacional_e_filtrado_pela_uf(tmp_path: Path) -> None:
    zp = tmp_path / "eleitorado_local_votacao_2024.zip"
    linhas = [_linha(GERACAO["antes"], ANTES[0]), _linha(GERACAO["antes"], (1, SAO, 1, 5, 1015, 1015, 50, None, None),
                                                          uf="SP")]
    with zipfile.ZipFile(zp, "w") as zf:
        zf.writestr("eleitorado_local_votacao_2024.csv", _csv(EL_HEADER, linhas))
    rj = dc.ler_csv(zp, ["RJ"])
    assert rj.select("SG_UF", "NR_SECAO").rows() == [("RJ", 10)] and rj.schema["NR_SECAO"] == pl.Int64
    assert "DT_GERACAO" not in rj.columns and rj.schema["QT_ELEITOR_SECAO"] == pl.String
    assert dc.comparar_zips(zp, zp, ["RJ", "SP"]).celulas.is_empty()  # um CSV para as duas UFs: tudo de uma vez


def test_chave_repetida_e_erro() -> None:
    df = pl.DataFrame({c: [1, 1] for c in dc.INTEIROS}).with_columns(pl.lit("RJ").alias("SG_UF"))
    with pytest.raises(v.TseDataError, match=r"antes: 1 linha\(s\) com a chave repetida"):
        dc.comparar_linhas(df, df)
    with pytest.raises(v.TseDataError, match="colunas da chave: NR_SECAO"):
        dc.comparar_linhas(df.drop("NR_SECAO"), df.drop("NR_SECAO"))


def test_cli_compara_com_a_versao_do_cache(acervo: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    saida = tmp_path / "dif.xlsx"
    assert cli.main(["--ano", "2026", "--diferencas", "2026-10-16", "--uf", "RJ", "--cache-dir", str(acervo),
                     "--saida", str(saida)]) == 0
    out = capsys.readouterr().out
    assert f"antes:  {CARIMBO['antes']} (gerada pelo TSE em 16/10/2026 06:29:01)" in out
    assert f"depois: {CARIMBO['depois']}" in out
    assert "Linhas (seção × turno): 10 → 10 · 2 criadas · 2 extintas" in out
    assert "Células alteradas: 8, em 4 linhas" in out
    assert "1º turno 1 → 1 · 2º turno 0 → 1" in out
    assert "VOLTA_AO_ORIGINAL → —: 1 seções, 280 eleitores (RJ)" in out
    assert f'DS_ENDERECO: "RUA DO LAVRADIO, 12" → "{ACESSO}" (seções 12)' in out
    wb = openpyxl.load_workbook(saida, read_only=True)
    assert wb.sheetnames == ["Resumo", "UFs", "Alterações", "Células", "Criadas e extintas", "Situações",
                             "Mudou de situação", "Consulta"]
    consulta = dict(list(wb["Consulta"].iter_rows(values_only=True))[1:])
    assert consulta["antes: versão do arquivo (carimbo)"] == CARIMBO["antes"]
    assert consulta["depois: versão do arquivo (carimbo)"] == CARIMBO["depois"]
    assert consulta["células alteradas"] == "8" and consulta["seções que mudaram de situação ou de local no 2º turno"] == "2"
    # a versão do cache ficou guardada: sem versões, as duas mais recentes são as mesmas
    assert [x.carimbo for x in vs.listar(acervo, SPEC.zip_name)] == [CARIMBO["antes"], CARIMBO["depois"]]
    assert cli.main(["--ano", "2026", "--diferencas", "--uf", "RJ", "--cache-dir", str(acervo),
                     "--saida", str(tmp_path / "dif2.xlsx")]) == 0
    assert "Células alteradas: 8, em 4 linhas" in capsys.readouterr().out


def test_destino_padrao_por_ufs_e_versoes(versoes: tuple[vs.Versao, vs.Versao]) -> None:
    nome = f"diferencas_cadastro_2026_{CARIMBO['antes']}_{CARIMBO['depois']}.xlsx"
    assert cli.destino_diferencas(["RJ"], 2026, *versoes) == Path("saidas/RJ") / nome
    assert cli.destino_diferencas(["RJ", "SP"], 2026, *versoes) == Path("saidas/RJ_SP") / nome
    assert cli.destino_diferencas(list(cli.uf_mod.UFS), 2026, *versoes) == Path("saidas/BR") / nome


def test_cli_erros_de_versao(tmp_path: Path, caplog: pytest.LogCaptureFixture) -> None:
    cache = escrever(tmp_path / "cache", "antes")
    base = ["--ano", "2026", "--uf", "RJ", "--cache-dir", str(cache)]
    assert cli.main([*base, "--diferencas"]) == 1
    assert "há 1 guardada(s)" in caplog.text
    assert cli.main([*base, "--diferencas", "2026-10-16"]) == 1
    assert "as duas versões são a mesma (2026-10-16T062901)" in caplog.text
    assert cli.main([*base, "--diferencas", "1999"]) == 1
    assert "nenhuma versão guardada com '1999'" in caplog.text
    with pytest.raises(SystemExit):
        cli.main([*base, "--diferencas", "a", "b", "c"])


@pytest.mark.dados
def test_golden_rj_2026_de_15_a_17_de_outubro() -> None:
    """RJ: de 15/10 para 16/10 nada mudou; de 16/10 para 17/10, só o endereço do local 1198 da zona 157 de Nova
    Iguaçu no 2º turno (5 seções, endereço e endereço original): "Rua Itabuna, 220" → "Rua do Imperador, 08 (acesso
    pela Rua Itabuna, 220)"."""
    cache, nome = Path("cache_tse"), SPEC.zip_name
    try:
        x = {c: vs.escolher(cache, nome, c) for c in ("2026-10-15", "2026-10-16", "2026-10-17")}
    except v.TseDataError:
        pytest.skip("acervo do cadastro de 2026 (15/10, 16/10 e 17/10) fora do cache")
    sem = dc.comparar_zips(x["2026-10-15"].zip, x["2026-10-16"].zip, ["RJ"])
    assert sem.celulas.is_empty() and sem.criadas.is_empty() and sem.extintas.is_empty()
    d = dc.comparar_zips(x["2026-10-16"].zip, x["2026-10-17"].zip, ["RJ"])
    assert d.linhas_antes == d.linhas_depois == 77_476 and d.criadas.is_empty() and d.extintas.is_empty()
    alt = dc.alteracoes(d)
    assert alt.select("NM_MUNICIPIO", "NR_ZONA", "NR_LOCAL_VOTACAO", "NR_TURNO", "COLUNA", "SECOES").rows() == [
        ("NOVA IGUAÇU", 157, 1198, 2, "DS_ENDERECO", "113, 114, 115, 116, 117"),
        ("NOVA IGUAÇU", 157, 1198, 2, "DS_ENDERECO_LOCVT_ORIGINAL", "113, 114, 115, 116, 117")]
    assert set(alt["DEPOIS"]) == {"RUA DO IMPERADOR, 08 (ACESSO PELA RUA ITABUNA, 220)"}
