"""Planilha por candidato: totais, eleitorado 2024/2026, destaque e abas — offline."""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import polars as pl
import pytest

import planilha_candidato as pc_cli
import votos_por_local_votacao as v
from conftest import CAND

RIO_ARGS = ["--ano", "2024", "--uf", "RJ", "--cargo", "vereador", "--candidato", str(CAND),
            "--municipio", "Rio de Janeiro"]


def _sheet(wb: openpyxl.Workbook, name: str) -> pl.DataFrame:
    rows = list(wb[name].values)
    return pl.DataFrame(rows[1:], schema=list(rows[0]), orient="row", infer_schema_length=None)


@pytest.fixture()
def workbook(tse_cache: Path, tmp_path: Path) -> openpyxl.Workbook:
    out = tmp_path / "p.xlsx"
    assert pc_cli.main(RIO_ARGS + ["--cache-dir", str(tse_cache), "--saida", str(out)]) == 0
    return openpyxl.load_workbook(out)


def test_sheets_and_totals(workbook: openpyxl.Workbook) -> None:
    assert workbook.sheetnames == ["Resumo", "Por local", "Por zona", "Por bairro", "Por secao",
                                   "Mudancas locais", "Mudancas secoes", "Inconsistencias", "Proveniencia"]
    for name in ("Por local", "Por zona", "Por bairro", "Por secao"):
        df = _sheet(workbook, name)
        assert df["VOTOS_CANDIDATO"].sum() == 29, name  # Niterói (mesmo nº, outra pessoa) fora
        assert df["VOTOS_VALIDOS"].sum() == 29 + 155 + 5, name
    resumo = dict(workbook["Resumo"].iter_rows(min_row=2, values_only=True))
    assert resumo["Votos do candidato"] == 29 and isinstance(resumo["Brancos"], int)
    assert resumo["Candidato"] == f"{CAND} — FULANA DE TAL"


def test_electorate_2024_2026(workbook: openpyxl.Workbook) -> None:
    loc = {r["NR_LOCAL_VOTACAO"]: r for r in _sheet(workbook, "Por local").iter_rows(named=True)}
    assert (loc[1015]["ELEITORADO_2024"], loc[1015]["ELEITORADO_2026"]) == (920, 955)
    assert (loc[1023]["ELEITORADO_2024"], loc[1023]["ELEITORADO_2026"]) == (350, 150)
    assert loc[1031]["ELEITORADO_2026"] is None
    zona = _sheet(workbook, "Por zona").row(0, named=True)
    assert (zona["ELEITORADO_2024"], zona["ELEITORADO_2026"]) == (1370, 1365)  # zona inteira, 1040 incluído
    bairro = {r["BAIRRO"]: r for r in _sheet(workbook, "Por bairro").iter_rows(named=True)}
    assert bairro["CENTRO"]["VOTOS_CANDIDATO"] == 14 and bairro["LAPA"]["VOTOS_CANDIDATO"] == 15
    assert (bairro["CENTRO"]["ELEITORADO_2024"], bairro["CENTRO"]["ELEITORADO_2026"]) == (1020, 1105)
    assert (bairro["LAPA"]["ELEITORADO_2024"], bairro["LAPA"]["ELEITORADO_2026"]) == (350, 260)


def test_change_columns_and_highlight(workbook: openpyxl.Workbook) -> None:
    df = _sheet(workbook, "Por local")
    loc = {r["NR_LOCAL_VOTACAO"]: r for r in df.iter_rows(named=True)}
    assert loc[1015]["MUDANCA_2024_2026"] == "MANTIDO" and loc[1015]["LOCAL_2026"] is None
    assert "RENOMEADO" in loc[1023]["MUDANCA_2024_2026"]
    assert loc[1023]["LOCAL_2026"].startswith("ESCOLA X RENOVADA")
    assert loc[1031]["MUDANCA_2024_2026"] == "DESATIVADO_EM_2026"

    ws = workbook["Por local"]
    col = openpyxl.utils.get_column_letter(df.columns.index("MUDANCA_2024_2026") + 1)
    rules = [(str(cf.sqref), rule.formula[0]) for cf in ws.conditional_formatting for rule in cf.rules]
    assert any(f"${col}2" in f and "DESATIVADO" in f for _, f in rules)  # vermelho
    assert any(f'${col}2<>"MANTIDO"' in f for _, f in rules)  # amarelo
    assert all(s.startswith("A2:") for s, _ in rules)  # linha inteira

    sec = {r["NR_SECAO"]: r["STATUS_SECAO"] for r in _sheet(workbook, "Por secao").iter_rows(named=True)}
    assert sec[21] == "MUDOU_DE_LOCAL" and sec[30] == "SECAO_EXTINTA_EM_2026"
    mud = _sheet(workbook, "Mudancas locais")
    assert 1040 in mud["NR_LOCAL_VOTACAO"].to_list()  # local novo aparece mesmo sem votos em 2024


def test_inconsistencies_sheet(workbook: openpyxl.Workbook) -> None:
    tipos = set(_sheet(workbook, "Inconsistencias")["TIPO"].to_list())
    assert {"BAIRRO_DIVERGENTE", "VARIACAO_ELEITORADO", "COORDENADA_AUSENTE",
            "SECOES_DIVERGEM_DA_LISTA_TRE", "LOCAL_2026_AUSENTE_NA_LISTA_TRE"} <= tipos
    prov = dict(workbook["Proveniencia"].iter_rows(min_row=2, values_only=True))
    assert prov["config.candidate"] == str(CAND)


def test_ambiguous_number_requires_municipality(tse_cache: Path) -> None:
    args = [a for a in RIO_ARGS if a not in ("--municipio", "Rio de Janeiro")]
    assert pc_cli.main(args + ["--cache-dir", str(tse_cache)]) == 1


@pytest.fixture()
def cache_2022(tse_cache: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Resultado de 2022 (cópia dos votos de 2024), sem rede: só vale o que estiver no cache."""
    shutil.copy(tse_cache / "votacao_secao_2024_RJ.zip", tse_cache / "votacao_secao_2022_RJ.zip")
    real = v.download

    def offline(spec: v.DatasetSpec, cache_dir: Path, verify: bool = False) -> Path:
        if not (cache_dir / spec.zip_name).exists():
            raise v.TseDataError(f"sem rede no teste: {spec.zip_name}")
        return real(spec, cache_dir, verify)

    monkeypatch.setattr(v, "download", offline)
    return tse_cache


def _run_2022(cache: Path, out: Path, *extra: str) -> int:
    return pc_cli.main(["--ano", "2022"] + RIO_ARGS[2:] + ["--cache-dir", str(cache), "--saida", str(out), *extra])


def test_base_year_defaults_to_result_year(cache_2022: Path, tmp_path: Path) -> None:
    """--ano 2022 compara o cadastro de 2022 com o de 2026 e traz eleitorado 2022, 2024 e 2026."""
    shutil.copy(cache_2022 / "eleitorado_local_votacao_2024.zip", cache_2022 / "eleitorado_local_votacao_2022.zip")
    out = tmp_path / "p22.xlsx"
    assert _run_2022(cache_2022, out) == 0
    wb = openpyxl.load_workbook(out)
    loc = {r["NR_LOCAL_VOTACAO"]: r for r in _sheet(wb, "Por local").iter_rows(named=True)}
    assert "MUDANCA_2024_2026" not in next(iter(loc.values()))
    assert loc[1015]["MUDANCA_2022_2026"] == "MANTIDO"
    assert loc[1031]["MUDANCA_2022_2026"] == "DESATIVADO_EM_2026"
    assert (loc[1015]["ELEITORADO_2022"], loc[1015]["ELEITORADO_2024"], loc[1015]["ELEITORADO_2026"]) == (920, 920, 955)
    assert loc[1015]["VAR_ELEITORADO"] == 35  # de 2022 para 2026
    assert "LOCAL_2022" in _sheet(wb, "Por secao").columns
    assert "VALOR_2022" in _sheet(wb, "Inconsistencias").columns
    resumo = dict(wb["Resumo"].iter_rows(min_row=2, values_only=True))
    assert resumo["Comparação de locais"] == "cadastro 2022 × cadastro 2026"
    assert resumo["Locais do resultado (2022) com mudança 2022→2026"] == 2


def test_base_year_register_is_required(cache_2022: Path, tmp_path: Path) -> None:
    assert _run_2022(cache_2022, tmp_path / "x.xlsx") == 1  # sem cadastro de 2022 no cache
    assert _run_2022(cache_2022, tmp_path / "x.xlsx", "--comparar-com", "2026") == 1


def test_compare_with_other_year(cache_2022: Path, tmp_path: Path) -> None:
    """--comparar-com 2024 num resultado de 2022 sem cadastro próprio: bairro cai para 2026/2024."""
    out = tmp_path / "p22.xlsx"
    assert _run_2022(cache_2022, out, "--comparar-com", "2024") == 0
    loc = _sheet(openpyxl.load_workbook(out), "Por local")
    assert "MUDANCA_2024_2026" in loc.columns and "ELEITORADO_2022" not in loc.columns
    assert set(loc["FONTE_CADASTRO"].to_list()) == {"2026", "2024"}  # 1031 só existe em 2024
