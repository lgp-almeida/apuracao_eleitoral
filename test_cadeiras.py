"""Distribuição das cadeiras de deputado: quocientes, sobras em 2 fases, entradas e API — offline."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import polars as pl
import pytest
from fastapi.testclient import TestClient

from apuracao import cadeiras as cd
from apuracao.divulgacao.cliente import ClienteDivulgacao, LimitadorTaxa
from apuracao.divulgacao.coletor import Coletor
from apuracao.web.app import create_app
from conftest import FakeTSE, _csv, download_sem_rede


# --------------------------------------------------------------------------- quociente eleitoral
@pytest.mark.parametrize("validos,vagas,qe", [
    (5_781_049, 55, 105_110),   # exemplo do documento: 105.109,98 → 105.110 (fração > 0,5 sobe)
    (8_395_113, 70, 119_930),   # Dep. Estadual RJ 2022 (oficial)
    (8_575_988, 46, 186_435),   # Dep. Federal RJ 2022 (oficial)
    (11, 2, 5),                 # 5,5: fração igual a 0,5 é desprezada
    (17, 3, 6),                 # 5,67: fração > 0,5 sobe
])
def test_quociente_eleitoral(validos: int, vagas: int, qe: int) -> None:
    assert cd.quociente_eleitoral(validos, vagas) == qe


# --------------------------------------------------------------------------- cenário à mão
# 10 vagas, 1.000 válidos → QE 100; 10% = 10, 20% = 20, 80% = 80.
AGREMIACOES = pl.DataFrame({"AGREMIACAO": ["A", "B", "C", "D", "E"], "NOME": ["A", "B", "C", "D", "E"],
                            "VOTOS": [450, 260, 170, 85, 35]})
CANDIDATOS = pl.DataFrame([
    # A: QP 4. A4 e A5 empatam (15): A5 tem desempate menor (mais idoso / seq do TSE) e fica com a vaga;
    # nenhum dos dois tem 20% do QE, então A não entra na fase 1 das sobras
    ("A", 11, "A1", 200, True, 1), ("A", 12, "A2", 150, True, 2), ("A", 13, "A3", 50, True, 3),
    ("A", 14, "A4", 15, True, 9), ("A", 15, "A5", 15, True, 1), ("A", 16, "A6", 5, True, 6),
    # B: QP 2, mas só B1 tem 10% do QE → a 2ª vaga do QP vai para as sobras
    ("B", 21, "B1", 240, True, 1), ("B", 22, "B2", 9, True, 2), ("B", 23, "B3", 5, True, 3),
    # C: QP 1; C4 é sub judice (voto não válido): não pode assumir, apesar dos 500 votos
    ("C", 31, "C1", 90, True, 1), ("C", 32, "C2", 60, True, 2), ("C", 33, "C3", 20, True, 3),
    ("C", 34, "C4", 500, False, 4),
    # D: QP 0, 85 votos (≥ 80% do QE); D2 com 15 (< 20% do QE) não entra na fase 1
    ("D", 41, "D1", 50, True, 1), ("D", 42, "D2", 15, True, 2),
    # E: 35 votos, abaixo de 80% do QE
    ("E", 51, "E1", 35, True, 1),
], schema=["AGREMIACAO", "NUMERO", "NOME", "VOTOS", "VALIDO", "DESEMPATE"], orient="row")


def test_distribuicao_calculada_a_mao() -> None:
    d = cd.distribuir(AGREMIACOES, CANDIDATOS, vagas=10, validos=1000)
    assert d.qe == 100 and d.vagas_nao_preenchidas == 0
    sit = dict(zip(d.candidatos["NOME"], d.candidatos["SITUACAO_PROJETADA"]))
    assert {n for n, s in sit.items() if s == cd.ELEITO_QP} == {"A1", "A2", "A3", "A5", "B1", "C1"}
    assert {n for n, s in sit.items() if s == cd.ELEITO_MEDIA} == {"C2", "D1", "C3", "B2"}
    assert sit["A4"] == cd.SUPLENTE and sit["C4"] == cd.NAO_ELEITO and sit["E1"] == cd.SUPLENTE
    # sobras: C e D empatam na média (85) → vence C, mais votada; fase 2 abre para todas → B (130)
    assert [(s["FASE"], s["AGREMIACAO"], s["NOME"]) for s in d.sobras] == [
        (1, "C", "C2"), (1, "D", "D1"), (1, "C", "C3"), (2, "B", "B2")]
    assert d.sobras[0]["MEDIA"] == 85 and d.sobras[3]["MEDIA"] == 130
    ag = {r["AGREMIACAO"]: r for r in d.agremiacoes.iter_rows(named=True)}
    assert {a: r["VAGAS"] for a, r in ag.items()} == {"A": 4, "B": 2, "C": 3, "D": 1, "E": 0}
    assert (ag["B"]["QP"], ag["B"]["VAGAS_QP"], ag["B"]["VAGAS_MEDIA"]) == (2, 1, 1)
    margem = dict(zip(d.candidatos["NOME"], d.candidatos["MARGEM"]))
    assert margem["A1"] == 200 - 15      # à frente do 1º suplente de A (A4)
    assert margem["A4"] == 15 - 15 + 1   # falta 1 voto para passar o último eleito de A (A5)
    assert margem["D2"] == 50 - 15 + 1 and margem["E1"] is None  # E não tem cadeira
    assert d.limites == {"cand_qp": 10.0, "cand_sobras": 20.0, "agrem_sobras": 80.0}


def test_vaga_sem_candidato_apto_fica_vazia() -> None:
    agr = pl.DataFrame({"AGREMIACAO": ["A"], "NOME": ["A"], "VOTOS": [100]})
    cand = pl.DataFrame({"AGREMIACAO": ["A"], "NUMERO": [1], "NOME": ["A1"], "VOTOS": [100], "VALIDO": [True]})
    d = cd.distribuir(agr, cand, vagas=3, validos=100)
    assert d.vagas_nao_preenchidas == 2 and d.agremiacoes["VAGAS"].to_list() == [1]


def test_sem_desempate_vence_o_menor_numero() -> None:
    cand = CANDIDATOS.drop("DESEMPATE")
    sit = dict(zip(*cd.distribuir(AGREMIACOES, cand, 10, 1000).candidatos.select("NOME", "SITUACAO_PROJETADA")))
    assert sit["A4"] == cd.ELEITO_QP and sit["A5"] == cd.SUPLENTE  # 14 < 15


# --------------------------------------------------------------------------- entradas
def test_entrada_divulgacao_federacao_e_sub_judice() -> None:
    base = {"CARGO": 7, "ABRANGENCIA": "uf", "UF": "RJ"}
    totais = pl.DataFrame([{**base, "VAGAS": 2, "VALIDOS": 300}])
    partidos = pl.DataFrame([{**base, "PARTIDO": "PX", "FEDERACAO": "FED", "VOTOS_TOTAL": 100},
                             {**base, "PARTIDO": "PY", "FEDERACAO": "FED", "VOTOS_TOTAL": 100},
                             {**base, "PARTIDO": "PZ", "FEDERACAO": None, "VOTOS_TOTAL": 100}])
    cand = pl.DataFrame([
        {**base, "NUMERO": 1, "NOME_URNA": "X", "PARTIDO": "PX", "FEDERACAO": "FED", "VOTOS": 90, "DESTINACAO": "Válido",
         "SEQ": 2, "SITUACAO": "Eleito por QP"},
        {**base, "NUMERO": 2, "NOME_URNA": "Y", "PARTIDO": "PY", "FEDERACAO": "FED", "VOTOS": 80, "DESTINACAO": "Válido",
         "SEQ": 3, "SITUACAO": "Suplente"},
        {**base, "NUMERO": 3, "NOME_URNA": "Z", "PARTIDO": "PZ", "FEDERACAO": None, "VOTOS": 999,
         "DESTINACAO": "Anulado sub judice", "SEQ": 1, "SITUACAO": "Não eleito"},
    ])
    agr, c, vagas, validos = cd.entrada_divulgacao(totais, cand, partidos, 7, "rj")
    assert (vagas, validos) == (2, 300)
    assert dict(zip(agr["AGREMIACAO"], agr["VOTOS"])) == {"FED": 200, "PZ": 100}  # federação = um só partido
    assert agr.filter(pl.col("AGREMIACAO") == "FED")["NOME"][0] == "FED (PX/PY)"
    assert c.filter(pl.col("NUMERO") == 3)["VALIDO"][0] is False and c["DESEMPATE"].to_list() == [2, 3, 1]
    with pytest.raises(cd.v.TseDataError):
        cd.entrada_divulgacao(totais, cand, partidos, 6, "RJ")


def _munzona(cache: Path) -> None:
    """votacao_partido/candidato_munzona sintéticos (2 zonas, federação, anulado, suplementar descartada)."""
    ph = ["CD_TIPO_ELEICAO", "NR_TURNO", "CD_CARGO", "SG_PARTIDO", "NR_FEDERACAO", "SG_FEDERACAO",
          "QT_TOTAL_VOTOS_LEG_VALIDOS", "QT_VOTOS_NOMINAIS_VALIDOS"]
    prow = [[2, 1, 7, "PX", 999, "FED", 10, 150], [2, 1, 7, "PX", 999, "FED", 0, 40],  # 2 zonas
            [2, 1, 7, "PY", 999, "FED", 0, 0], [2, 1, 7, "PZ", -1, "#NULO#", 5, 95],
            [1, 1, 7, "PZ", -1, "#NULO#", 0, 5000]]  # eleição suplementar: fora
    ch = ["CD_TIPO_ELEICAO", "NR_TURNO", "CD_CARGO", "SQ_CANDIDATO", "NR_CANDIDATO", "NM_URNA_CANDIDATO", "SG_PARTIDO",
          "NR_FEDERACAO", "SG_FEDERACAO", "NM_TIPO_DESTINACAO_VOTOS", "QT_VOTOS_NOMINAIS_VALIDOS", "DS_SIT_TOT_TURNO"]
    crow = [[2, 1, 7, 1, 11, "X1", "PX", 999, "FED", "Válido", 120, "ELEITO POR QP"],
            [2, 1, 7, 1, 11, "X1", "PX", 999, "FED", "Válido", 40, "ELEITO POR QP"],
            [2, 1, 7, 2, 12, "X2", "PX", 999, "FED", "Válido", 30, "ELEITO POR MÉDIA"],
            [2, 1, 7, 3, 31, "Z1", "PZ", -1, "#NULO#", "Válido", 95, "ELEITO POR MÉDIA"],
            [2, 1, 7, 4, 32, "Z2", "PZ", -1, "#NULO#", "Anulado", 0, "NÃO ELEITO"]]
    for nome, h, rows in (("votacao_partido_munzona_2022", ph, prow), ("votacao_candidato_munzona_2022", ch, crow)):
        with zipfile.ZipFile(cache / f"{nome}.zip", "w") as zf:
            zf.writestr(f"{nome}_RJ.csv", _csv(h, rows))


def test_entrada_munzona_reproduz_o_oficial(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    download_sem_rede(monkeypatch)
    _munzona(tmp_path)
    agr, cand, vagas, validos = cd.entrada_munzona(2022, "RJ", 7, tmp_path)
    assert dict(zip(agr["AGREMIACAO"], agr["VOTOS"])) == {"FED": 200, "PZ": 100}  # zonas somadas; suplementar fora
    assert (vagas, validos) == (3, 300)  # vagas = eleitos oficiais
    assert dict(zip(cand["NUMERO"], cand["VOTOS"]))[11] == 160 and not cand.filter(pl.col("NUMERO") == 32)["VALIDO"][0]
    d = cd.distribuir(agr, cand, vagas, validos)  # QE 100: FED QP 2 (X1; X2 com 30 ≥ 10), PZ QP 1
    # X2 (30 ≥ 10% do QE) e Z1 entram pelo QP no cálculo; o "oficial" sintético diz média: as 2 são apontadas
    assert set(cd.comparar_com_oficial(d)["NUMERO"]) == {12, 31}
    assert set(d.candidatos.filter(pl.col("SITUACAO_PROJETADA").str.starts_with("Eleito"))["NUMERO"]) == {11, 12, 31}


# --------------------------------------------------------------------------- API
@pytest.fixture()
def site(fake_tse: FakeTSE, tse_cache: Path, tmp_path: Path) -> TestClient:
    dados = tmp_path / "dados"
    Coletor(ClienteDivulgacao("simulado", sessao=fake_tse, limitador=LimitadorTaxa(1e9)), dados).ciclo()
    return TestClient(create_app(dados, "RJ", tse_cache))


def test_api_cadeiras(site: TestClient) -> None:
    r = site.get("/api/cadeiras?cargo=7").json()
    assert r["fonte"] == "divulgação do TSE" and r["qe"] == cd.quociente_eleitoral(r["validos"], r["vagas"])
    assert sum(a["VAGAS"] for a in r["agremiacoes"]) == r["vagas"] - r["vagas_nao_preenchidas"]
    assert len(r["eleitos"]) == r["eleitos_qp"] + r["eleitos_media"] == r["vagas"] - r["vagas_nao_preenchidas"]
    # o recorte do simulado nos testes é truncado (4 agremiações, válidos do estado inteiro): a
    # salvaguarda precisa apontar isso (com os dados completos do simulado, soma = válidos e 0 divergências)
    assert r["consistente"] is False and r["soma_agremiacoes"] < r["validos"]
    assert site.get("/api/cadeiras?cargo=3").status_code == 400
    card = next(c for c in site.get("/api/painel").json()["cartoes"] if c["cargo"] == 7)
    assert card["cadeiras"]["vagas"] == r["vagas"] and "eleitos" not in card["cadeiras"]
    eleito = r["eleitos"][0]
    k = site.get(f"/api/candidato?cargo=7&numero={eleito['NUMERO']}").json()["cadeira"]
    assert k["situacao"] == eleito["SITUACAO_PROJETADA"] and k["qe"] == r["qe"]


def test_api_cadeiras_historico_usa_microdados(tmp_path: Path, tse_cache: Path, site: TestClient,
                                               monkeypatch: pytest.MonkeyPatch) -> None:
    download_sem_rede(monkeypatch)
    _munzona(tse_cache)
    dados = tmp_path / "hist"
    (dados / "ultimo").mkdir(parents=True)
    (dados / "status.json").write_text(json.dumps({"ano": 2022, "turno": 1, "uf": "RJ"}))
    r = TestClient(create_app(dados, "RJ", tse_cache)).get("/api/cadeiras?cargo=7").json()
    assert r["fonte"] == "microdados oficiais do TSE (2022)" and r["vagas"] == 3 and r["qe"] == 100
    assert r["conferencia_tse"] == {"eleitos_tse": 3, "coincidentes": 3, "divergencias": 0}


# --------------------------------------------------------------------------- dados reais (regressão)
CACHE_REAL = Path(__file__).parent / "cache_tse"


@pytest.mark.skipif(not (CACHE_REAL / "votacao_candidato_munzona_2022.zip").exists()
                    or not (CACHE_REAL / "votacao_partido_munzona_2022.zip").exists(),
                    reason="microdados oficiais de 2022 fora do cache")
@pytest.mark.parametrize("cargo,vagas,qe,qp,media", [(6, 46, 186_435, 36, 10), (7, 70, 119_930, 58, 12)])
def test_reproduz_os_eleitos_do_rj_em_2022(cargo: int, vagas: int, qe: int, qp: int, media: int) -> None:
    """Os 116 deputados do RJ em 2022, com QP × média, iguais ao resultado oficial (sem divergência)."""
    d = cd.distribuir(*cd.entrada_munzona(2022, "RJ", cargo, CACHE_REAL))
    sit = d.candidatos["SITUACAO_PROJETADA"]
    assert (d.vagas, d.qe, (sit == cd.ELEITO_QP).sum(), (sit == cd.ELEITO_MEDIA).sum()) == (vagas, qe, qp, media)
    assert cd.comparar_com_oficial(d).is_empty()


def test_sem_votos_validos_nao_divide_por_zero() -> None:
    """Início da apuração (achado no ensaio geral): sem voto válido, QE seria 0 e a API dava 500."""
    with pytest.raises(cd.v.TseDataError):
        cd.distribuir(AGREMIACOES.with_columns(pl.lit(0).alias("VOTOS")), CANDIDATOS, vagas=10, validos=0)


def test_cadeiras_acompanham_cada_nova_coleta(site: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """O cache das cadeiras não pode usar id() das tabelas: o CPython reaproveita o endereço de um
    DataFrame liberado, e a chave de duas coletas atrás voltava a valer (resultado velho)."""
    import os

    import apuracao.web.app as appmod
    monkeypatch.setattr(appmod, "id", lambda obj: 42, raising=False)  # pior caso do reuso de endereço
    ultimo = tmp_path / "dados" / "ultimo"
    originais = {n: pl.read_parquet(ultimo / f"{n}.parquet") for n in ("totais", "candidatos", "partidos")}
    t0 = (ultimo / "partidos.parquet").stat().st_mtime_ns
    for g in range(1, 7):
        extra = 1000 * g
        part = originais["partidos"].with_columns(
            pl.when(pl.col("ABRANGENCIA") == "uf").then(pl.col("VOTOS_TOTAL") + extra).otherwise(pl.col("VOTOS_TOTAL")))
        for nome, df in (("totais", originais["totais"]), ("candidatos", originais["candidatos"]), ("partidos", part)):
            df.write_parquet(ultimo / f"{nome}.parquet")
            os.utime(ultimo / f"{nome}.parquet", ns=(t0 + g * 10**9, t0 + g * 10**9))
        esperado = int(part.filter((pl.col("CARGO") == 7) & (pl.col("ABRANGENCIA") == "uf"))["VOTOS_TOTAL"].sum())
        assert site.get("/api/cadeiras?cargo=7").json()["soma_agremiacoes"] == esperado, f"coleta {g}"


def test_planilha_das_listas_de_eleitos(site: TestClient) -> None:
    """Rodada 31: as listas de eleitos projetados em .xlsx, com a coluna "Destacado"."""
    import io
    import openpyxl
    r = site.get("/api/cadeiras?cargo=7").json()
    comp = {a["AGREMIACAO"]: a["PARTIDOS"] for a in r["composicao"]}
    assert set(comp) == {a["AGREMIACAO"] for a in r["agremiacoes"]}
    alvo = r["eleitos"][0]["PARTIDO"]
    resp = site.get(f"/api/cadeiras/planilha?cargo=7&destacar={alvo}")
    assert resp.status_code == 200 and "eleitos_projetados_deputado_estadual_RJ" in resp.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames[0] == "Sobre" and "Dep. Estadual RJ eleitos" in wb.sheetnames
    ws = wb["Dep. Estadual RJ eleitos"]
    cab = [c.value for c in ws[1]]
    linhas = [dict(zip(cab, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]
    assert len(linhas) == len(r["eleitos"])
    assert {str(x["Número"]) for x in linhas if x["Destacado"] == "Sim"} == \
        {str(e["NUMERO"]) for e in r["eleitos"] if e["PARTIDO"] == alvo}
    sobre = {row[0].value: row[1].value for row in wb["Sobre"].iter_rows(min_row=2)}
    assert sobre["Destaque"] == alvo and sobre["Cargo"].startswith("Deputado Estadual")
    sem = openpyxl.load_workbook(io.BytesIO(site.get("/api/cadeiras/planilha?cargo=7").content))["Dep. Estadual RJ eleitos"]
    assert all(c.value in (None, "") for c in list(sem.columns)[-1][1:])      # sem destaque: ninguém marcado
    assert site.get("/api/cadeiras/planilha?cargo=3").status_code == 400


def test_federacao_destaca_todos_os_seus_partidos() -> None:
    from apuracao.boletim import destacado
    fed = {"PARTIDO": "PC do B", "AGREMIACAO": "PT/PC do B/PV"}
    assert destacado(fed, frozenset({"PT/PC do B/PV"})) and destacado(fed, frozenset({"PC do B"}))
    assert not destacado(fed, frozenset({"PT"})) and not destacado(fed, frozenset())
