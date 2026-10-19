"""Local de votação de cada seção no 1º × no 2º turno (`apuracao/locais_turnos.py` + `locais_entre_turnos.py`).

Cenário sintético (Rio, zona 4; Niterói só para provar que a chave inclui o município), cadastro gerado pelo TSE
em 15/10/2026 14:33:26 com as linhas dos dois turnos:

  seção 10  1015 nos dois                                           → sem mudança
  seção 11  1º no 1050 (temporário; original 1015), 2º no 1015       → VOLTA_AO_ORIGINAL
  seção 12  1º no 1015, 2º no 1050 (temporário; original 1015)       → TEMPORARIO_NO_2T
  seção 13  agregada à 12: vai junto                                 → TEMPORARIO_NO_2T
  seção 20  1º no 1040 e 2º no 1050, os dois temporários do 1023     → OUTRO_TEMPORARIO
  seção 21  1º no 1023, 2º no 1040 que passa a ser o original        → OUTRO_LOCAL
  seção 22  só a pontuação do endereço muda                          → sem mudança
  seção 23  mesmo local, nome diferente no 2º                        → NOME_OU_ENDERECO
  seção 27  mesmo nº 1060, mas outro nome e outro endereço no 2º     → OUTRO_PREDIO
  seção 24  passa a ser agregada à 22 no 2º                          → AGREGACAO
  seção 25  só no 1º turno                                           → só na contagem
  seção 26  só no 2º turno                                           → SO_NO_2T
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import openpyxl
import polars as pl
import pytest

import locais_entre_turnos as cli
import votos_por_local_votacao as v
from apuracao import eleitorado as el
from apuracao import locais_turnos as lt
from apuracao import versoes as vs
from conftest import EL_HEADER, NIT, RIO, _csv

LOCAIS = {  # nº → (nome, endereço, bairro, latitude, longitude)
    1015: ("COLÉGIO PEDRO II", "AV. MARECHAL FLORIANO, 80", "CENTRO", "-22,9035", "-43,1790"),
    1023: ("ESCOLA X", "RUA DO LAVRADIO, 12", "LAPA", "-22,9130", "-43,1800"),
    1040: ("CRECHE NOVA", "RUA Z, 5", "LAPA", "-22,9200", "-43,1800"),
    1050: ("GINÁSIO PROVISÓRIO", "RUA W, 9", "CENTRO", "-22,9100", "-43,1700"),
    1060: ("CIEP 351 - GRAMA", "ESTRADA SANTA BARBARA S/N", "GRAMA", "-22,9000", "-43,1800"),
}
# (município, zona, seção, principal, local, local original, eleitores, outro nome, outro endereço)
TURNO_1 = [
    (RIO, 4, 10, -1, 1015, 1015, 300, None, None),
    (RIO, 4, 11, -1, 1050, 1015, 280, None, None),
    (RIO, 4, 12, -1, 1015, 1015, 250, None, None),
    (RIO, 4, 13, 12, 1015, 1015, 90, None, None),
    (RIO, 4, 20, -1, 1040, 1023, 200, None, None),
    (RIO, 4, 21, -1, 1023, 1023, 150, None, None),
    (RIO, 4, 22, -1, 1023, 1023, 120, None, None),
    (RIO, 4, 23, -1, 1023, 1023, 110, None, None),
    (RIO, 4, 24, -1, 1023, 1023, 100, None, None),
    (RIO, 4, 25, -1, 1023, 1023, 80, None, None),
    (RIO, 4, 27, -1, 1060, 1060, 60, None, None),
    (NIT, 71, 11, -1, 1015, 1015, 310, None, None),
]
TURNO_2 = [
    (RIO, 4, 10, -1, 1015, 1015, 300, None, None),
    (RIO, 4, 11, -1, 1015, 1015, 280, None, None),
    (RIO, 4, 12, -1, 1050, 1015, 250, None, None),
    (RIO, 4, 13, 12, 1050, 1015, 90, None, None),
    (RIO, 4, 20, -1, 1050, 1023, 200, None, None),
    (RIO, 4, 21, -1, 1040, 1040, 150, None, None),
    (RIO, 4, 22, -1, 1023, 1023, 120, None, "RUA DO LAVRADIO 12"),
    (RIO, 4, 23, -1, 1023, 1023, 110, "ESCOLA X RENOVADA", None),
    (RIO, 4, 24, 22, 1023, 1023, 100, None, None),
    (RIO, 4, 26, -1, 1015, 1015, 70, None, None),
    (RIO, 4, 27, -1, 1060, 1060, 60, "COLÉGIO ESTADUAL BARÃO DE TINGUÁ", "RUA GISELA, 157"),
    (NIT, 71, 11, -1, 1015, 1015, 310, None, None),
]
GERACAO = ("15/10/2026", "14:33:26")
CARIMBO = "2026-10-15T143326"


def _linha(turno: int, dado: tuple) -> list:
    (cd, nm), zona, secao, principal, local, original, eleitores, nome, endereco = dado
    nome_local, end_local, bairro, lat, lon = LOCAIS[local]
    campos = dict.fromkeys(EL_HEADER, "")
    campos.update(
        DT_GERACAO=GERACAO[0], HH_GERACAO=GERACAO[1], AA_ELEICAO=2026, NR_TURNO=turno, SG_UF="RJ",
        DT_ELEICAO="04/10/2026" if turno == 1 else "25/10/2026", DS_ELEICAO=f"{turno}º Turno",
        CD_MUNICIPIO=str(cd), NM_MUNICIPIO=nm, NR_ZONA=zona, NR_SECAO=secao, NR_SECAO_PRINCIPAL=principal,
        CD_TIPO_SECAO_AGREGADA=1 if principal == -1 else 2,
        DS_TIPO_SECAO_AGREGADA="Principal" if principal == -1 else "Agregada",
        NR_LOCAL_VOTACAO=local, NM_LOCAL_VOTACAO=nome or nome_local, DS_ENDERECO=endereco or end_local,
        NM_BAIRRO=bairro, NR_LATITUDE=lat, NR_LONGITUDE=lon, DS_SITU_LOCAL_VOTACAO="ATIVO",
        QT_ELEITOR_SECAO=eleitores, NR_LOCAL_VOTACAO_ORIGINAL=original,
        NM_LOCAL_VOTACAO_ORIGINAL=LOCAIS[original][0], DS_ENDERECO_LOCVT_ORIGINAL=LOCAIS[original][1],
    )
    return [campos[c] for c in EL_HEADER]


def escrever_cadastro(cache: Path, turno_2: list[tuple] = TURNO_2) -> Path:
    cache.mkdir(parents=True, exist_ok=True)
    linhas = [_linha(1, d) for d in TURNO_1] + [_linha(2, d) for d in turno_2]
    tmp = cache / "eleitorado_local_votacao_2026.zip.tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("eleitorado_local_votacao_2026_RJ.csv", _csv(EL_HEADER, linhas))
        zf.writestr("eleitorado_local_votacao_2026_BRASIL.csv", _csv(EL_HEADER, linhas))
        zf.writestr("leiame.pdf", b"%PDF-fake")
    tmp.replace(cache / "eleitorado_local_votacao_2026.zip")  # como o download: o ZIP guardado no acervo é só leitura
    return cache


@pytest.fixture
def cache_turnos(tmp_path: Path) -> Path:
    return escrever_cadastro(tmp_path / "cache")


@pytest.fixture
def comparacao(cache_turnos: Path) -> lt.Comparacao:
    return lt.comparar(el.load_sections(2026, "RJ", cache_turnos, turno=1),
                       el.load_sections(2026, "RJ", cache_turnos, turno=2))


def _situacoes(c: lt.Comparacao) -> dict[tuple[str, int], str]:
    return {(m, s): x for m, s, x in c.secoes.select("NM_MUNICIPIO", "NR_SECAO", "SITUACAO").iter_rows()}


def test_situacao_de_cada_secao(comparacao: lt.Comparacao) -> None:
    rio = "RIO DE JANEIRO"
    # 10, 22 (só pontuação), 25 (só no 1º turno) e Niterói ficam de fora
    assert _situacoes(comparacao) == {
        (rio, 11): "VOLTA_AO_ORIGINAL", (rio, 12): "TEMPORARIO_NO_2T", (rio, 13): "TEMPORARIO_NO_2T",
        (rio, 20): "OUTRO_TEMPORARIO", (rio, 21): "OUTRO_LOCAL", (rio, 23): "NOME_OU_ENDERECO",
        (rio, 24): "AGREGACAO", (rio, 26): "SO_NO_2T", (rio, 27): "OUTRO_PREDIO"}


def test_contagens_e_colunas_da_secao(comparacao: lt.Comparacao) -> None:
    assert comparacao.nos_dois_turnos == 11 and comparacao.so_no_1t == 1
    assert comparacao.eleitores_2t == sum(d[6] for d in TURNO_2)
    volta = comparacao.secoes.filter(pl.col("NR_SECAO") == 11).row(0, named=True)
    assert (volta["NR_LOCAL_VOTACAO_1T"], volta["NR_LOCAL_VOTACAO_2T"], volta["NR_LOCAL_VOTACAO_ORIGINAL"]) == \
        (1050, 1015, 1015)
    assert volta["NM_LOCAL_VOTACAO_ORIGINAL"] == "COLÉGIO PEDRO II" and 1000 < volta["DISTANCIA_M"] < 1300
    agregada = comparacao.secoes.filter(pl.col("NR_SECAO") == 13).row(0, named=True)
    assert agregada["TIPO_SECAO_2T"] == "AGREGADA" and agregada["NR_SECAO_PRINCIPAL_2T"] == 12
    nova = comparacao.secoes.filter(pl.col("NR_SECAO") == 26).row(0, named=True)
    assert nova["NR_LOCAL_VOTACAO_1T"] is None and nova["DISTANCIA_M"] is None
    assert comparacao.secoes.filter(pl.col("NR_SECAO") == 23)["DISTANCIA_M"].item() == 0


def test_resumo_tem_todas_as_situacoes(comparacao: lt.Comparacao, cache_turnos: Path) -> None:
    r = {x["SITUACAO"]: x for x in lt.resumo(comparacao).iter_rows(named=True)}
    assert list(r) == list(lt.SITUACOES)
    assert (r["TEMPORARIO_NO_2T"]["SECOES"], r["TEMPORARIO_NO_2T"]["ELEITORES"],
            r["TEMPORARIO_NO_2T"]["LOCAIS_NO_2T"]) == (2, 340, 1)
    turno_1 = el.load_sections(2026, "RJ", cache_turnos, turno=1)
    vazio = lt.resumo(lt.comparar(turno_1, turno_1))  # sem mudança: as situações aparecem com zero
    assert vazio.height == len(lt.SITUACOES) and vazio["SECOES"].sum() == vazio["ELEITORES"].sum() == 0


def test_locais_agrupa_as_secoes_de_cada_par(comparacao: lt.Comparacao) -> None:
    pares = lt.locais(comparacao)
    temp = pares.filter(pl.col("SITUACAO") == "TEMPORARIO_NO_2T").row(0, named=True)
    assert (temp["NR_LOCAL_VOTACAO_1T"], temp["NR_LOCAL_VOTACAO_2T"], temp["N_SECOES"], temp["SECOES"],
            temp["QT_ELEITORES"]) == (1015, 1050, 2, "12, 13", 340)
    assert pares.height == comparacao.secoes.height - 1  # 12 e 13 viram um par
    assert pares.filter(pl.col("SITUACAO") == "SO_NO_2T")["NR_LOCAL_VOTACAO_1T"].item() is None


def test_sem_as_colunas_do_local_original_toda_troca_e_outro_local(cache_turnos: Path) -> None:
    sem = [el.load_sections(2026, "RJ", cache_turnos, turno=t).drop(
        "NR_LOCAL_VOTACAO_ORIGINAL", "NM_LOCAL_VOTACAO_ORIGINAL", "DS_ENDERECO_LOCVT_ORIGINAL") for t in (1, 2)]
    sit = _situacoes(lt.comparar(*sem))
    assert {s for (_, n), s in sit.items() if n in (11, 12, 13, 20, 21)} == {"OUTRO_LOCAL"}


def test_cli_grava_planilha_com_a_versao_e_a_guarda_no_acervo(cache_turnos: Path, tmp_path: Path,
                                                              capsys: pytest.CaptureFixture[str]) -> None:
    saida = tmp_path / "turnos.xlsx"
    assert cli.main(["--cache-dir", str(cache_turnos), "--saida", str(saida)]) == 0
    out = capsys.readouterr().out
    assert f"versão {CARIMBO} (gerada pelo TSE em 15/10/2026 14:33:26)" in out
    assert "Zona 4 · 1050 GINÁSIO PROVISÓRIO → 1015 COLÉGIO PEDRO II [VOLTA_AO_ORIGINAL]" in out
    wb = openpyxl.load_workbook(saida)
    assert wb.sheetnames == ["Resumo", "Locais", "Seções", "Consulta"]
    consulta = {r[0]: r[1] for r in wb["Consulta"].iter_rows(min_row=2, values_only=True)}
    versao = vs.escolher(cache_turnos, "eleitorado_local_votacao_2026.zip", CARIMBO)
    assert consulta["versão do arquivo (carimbo)"] == CARIMBO and consulta["SHA-512"] == versao.sha512
    assert consulta["seções só no 1º turno"] == "1" and consulta["UFs"] == "RJ"
    assert versao.sha512 == json.loads((cache_turnos / "eleitorado_local_votacao_2026.proveniencia.json")
                                       .read_text())["sha512"]
    # a mesma comparação refeita com a versão guardada, mesmo depois de o cache trocar de versão
    escrever_cadastro(cache_turnos, TURNO_1)
    (cache_turnos / "eleitorado_local_votacao_2026.proveniencia.json").unlink()
    for p in cache_turnos.glob("eleitorado_local_votacao_2026__*.parquet"):
        p.unlink()
    refeita = tmp_path / "refeita.xlsx"
    assert cli.main(["--cache-dir", str(cache_turnos), "--versao", "2026-10-15", "--saida", str(refeita)]) == 0
    assert [r for r in openpyxl.load_workbook(refeita)["Seções"].iter_rows(values_only=True)] == \
        [r for r in wb["Seções"].iter_rows(values_only=True)]


def test_destino_padrao_leva_o_carimbo(cache_turnos: Path) -> None:
    v.download(v.electorate_spec(2026, "RJ"), cache_turnos)
    el.ensure_local_provenance(v.electorate_spec(2026, "RJ"), cache_turnos)
    versao = vs.em_uso(cache_turnos, "eleitorado_local_votacao_2026.zip")
    assert cli.destino_padrao(["RJ"], 2026, versao) == Path(f"saidas/RJ/locais_entre_turnos_2026_{CARIMBO}.xlsx")
    assert cli.destino_padrao(list(cli.uf_mod.UFS), 2026, versao).parent == Path("saidas/BR")


def test_cli_cadastro_sem_2o_turno_e_erro(tse_cache: Path, caplog: pytest.LogCaptureFixture) -> None:
    assert cli.main(["--cache-dir", str(tse_cache), "--saida", str(tse_cache / "x.xlsx")]) == 1
    assert "não tem linhas do 2º turno" in caplog.text


def test_cli_versao_inexistente_lista_as_guardadas(cache_turnos: Path, caplog: pytest.LogCaptureFixture) -> None:
    assert cli.main(["--cache-dir", str(cache_turnos), "--versao", "1999"]) == 1
    assert "nenhuma versão guardada com '1999'" in caplog.text


def test_cli_municipio_so_com_uma_uf(cache_turnos: Path) -> None:
    with pytest.raises(SystemExit):
        cli.main(["--cache-dir", str(cache_turnos), "--ufs", "RJ", "SP", "--municipio", "Niterói"])


def test_cli_filtra_o_municipio(cache_turnos: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert cli.main(["--cache-dir", str(cache_turnos), "--municipio", "Niterói", "--saida",
                     str(tmp_path / "nit.xlsx")]) == 0
    assert "Nenhuma seção muda de local entre os turnos." in capsys.readouterr().out


@pytest.mark.dados
def test_golden_rj_2026_versoes_do_cadastro() -> None:
    """RJ, as três versões do cadastro com os dois turnos. A de 06/10 repetia no 2º turno o local cadastrado: as 47
    seções em local temporário no 1º "voltavam". As de 15/10 e 16/10 deixam o 2º turno no mesmo local temporário; o
    que muda são 19 seções que mantêm o nº do local num prédio diferente (zona 156 de Nova Iguaçu, local 1279: CIEP
    351 → Colégio Estadual Barão de Tinguá; zona 188 do Rio, local 1295: EEEPV da Rua Muaná → da Rua Irapuá)."""
    cache, nome = Path("cache_tse"), "eleitorado_local_votacao_2026.zip"
    esperado = {"2026-10-06": {"VOLTA_AO_ORIGINAL": 47, "OUTRO_PREDIO": 14},
                "2026-10-15": {"OUTRO_PREDIO": 19}, "2026-10-16": {"OUTRO_PREDIO": 19}}
    try:
        versoes = {c: vs.escolher(cache, nome, c) for c in esperado}
    except v.TseDataError:
        pytest.skip("acervo do cadastro de 2026 (06/10, 15/10 e 16/10) fora do cache")
    for c, versao in versoes.items():
        comp = lt.comparar(*[el.load_sections(2026, "RJ", versao.pasta, turno=t) for t in (1, 2)])
        muda = comp.secoes.filter(pl.col("SITUACAO").is_in(lt.MUDA_DE_LOCAL))
        assert dict(muda.group_by("SITUACAO").len().iter_rows()) == esperado[c], c
        assert comp.nos_dois_turnos == 38_738 and comp.so_no_1t == 0
    predios = muda.filter(pl.col("SITUACAO") == "OUTRO_PREDIO").group_by("NR_ZONA", "NR_LOCAL_VOTACAO_1T").len()
    assert sorted(predios.rows()) == [(156, 1279, 14), (188, 1295, 5)]
