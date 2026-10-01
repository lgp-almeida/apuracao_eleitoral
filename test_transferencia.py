"""Transferência de votos do 1º para o 2º turno (apuracao/transferencia.py) — offline, dados sintéticos
com a matriz verdadeira conhecida; e regressão com 2022 real (se estiver no cache)."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import polars as pl
import pytest
from fastapi.testclient import TestClient

from apuracao import transferencia as tr
from apuracao.web.app import create_app

# 1º turno: A (10), B (20), C (30, 6%), D (40, 3%), E (50, 0,4% -> "Outros"), branco (95), nulo (96), abstenção
ORIGENS = [10, 20, 30, 40, 50, 95, 0]      # 0 = abstenção
# verdade por zona: linhas = origens acima; colunas = A, B, branco/nulo, abstenção
B_ZONA = {
    1: np.array([[.97, .01, .01, .01], [.01, .96, .01, .02], [.30, .50, .15, .05], [.20, .60, .15, .05],
                 [.50, .20, .20, .10], [.30, .20, .40, .10], [.05, .05, .00, .90]]),
    2: np.array([[.95, .02, .01, .02], [.02, .95, .01, .02], [.55, .25, .15, .05], [.40, .40, .10, .10],
                 [.50, .20, .20, .10], [.20, .30, .40, .10], [.03, .07, .00, .90]]),
}
NOMES = {10: "ANA (P1)", 20: "BIA (P2)", 30: "CAIO (P3)", 40: "DUDA (P4)", 50: "EVA (P5)"}


def sintetico(locais_por_zona: int = 60, secoes_por_local: int = 4, semente: int = 3
              ) -> tuple[pl.DataFrame, pl.DataFrame, dict[int, str]]:
    """Seções no formato de `carregar_microdados`; cada eleitor segue a matriz da sua zona (multinomial)."""
    rng = np.random.default_rng(semente)
    secoes, votos = [], []
    base = np.array([.33, .27, .06, .03, .004, .05, .256])
    for zona, bz in B_ZONA.items():
        for loc in range(locais_por_zona):
            perfil = rng.dirichlet(base * 60)            # locais diferentes entre si (identifica a matriz)
            for sec in range(secoes_por_local):
                aptos = int(rng.integers(250, 420))
                n1 = rng.multinomial(aptos, rng.dirichlet(perfil * 300))
                n2 = sum(rng.multinomial(n, bz[k]) for k, n in enumerate(n1))
                chave = {"CD_MUNICIPIO": 1000 + zona, "NR_ZONA": zona, "NR_SECAO": loc * 10 + sec}
                secoes.append({**chave, "NR_LOCAL_VOTACAO": loc, "NM_MUNICIPIO": f"CIDADE {zona}",
                               "NM_LOCAL_VOTACAO": f"ESCOLA {zona}-{loc}", "APTOS_1": aptos, "APTOS_2": aptos})
                for numero, q in zip(ORIGENS, n1):
                    if numero:
                        votos.append({"NR_TURNO": 1, **chave, "NR_VOTAVEL": numero, "QT_VOTOS": int(q)})
                for numero, q in zip([10, 20, 95], n2[:3]):
                    votos.append({"NR_TURNO": 2, **chave, "NR_VOTAVEL": numero, "QT_VOTOS": int(q)})
    return pl.DataFrame(secoes), pl.DataFrame(votos), dict(NOMES)


@pytest.fixture(scope="module")
def dados():
    return sintetico()


# --------------------------------------------------------------------------- estimação
def test_projecao_no_simplex() -> None:
    m = np.array([[0.2, 0.3, 0.5], [2.0, -1.0, 0.0], [-3.0, -3.0, -3.0], [0.9, 0.9, 0.9]])
    p = tr.projetar_simplex(m)
    assert np.allclose(p.sum(axis=1), 1) and (p >= 0).all()
    assert np.allclose(p[0], m[0])                          # já no simplex: não muda
    assert np.allclose(p[1], [1, 0, 0]) and np.allclose(p[2], [1 / 3] * 3)


def test_recupera_a_matriz_verdadeira() -> None:
    rng = np.random.default_rng(1)
    verdade = B_ZONA[1][[0, 1, 2, 5, 6]]
    x = rng.dirichlet(np.ones(5) * 2, 3000)
    y = x @ verdade + rng.normal(0, 0.003, (3000, 4))
    b = tr.estimar(x, y, np.ones(3000))
    assert np.abs(b - verdade).max() < 0.02
    assert np.allclose(b.sum(axis=1), 1) and (b >= 0).all()


def test_lote_igual_a_um_por_vez() -> None:
    rng = np.random.default_rng(2)
    gs, cs = [], []
    for _ in range(3):
        x = rng.dirichlet(np.ones(4), 200)
        y = x @ rng.dirichlet(np.ones(3), 4)
        gs.append(x.T @ x)
        cs.append(x.T @ y)
    lote = tr._fista_lote(np.stack(gs), np.stack(cs))
    for k in range(3):
        assert np.abs(lote[k] - tr._fista(gs[k], cs[k])).max() < 1e-8


# --------------------------------------------------------------------------- unidades
def test_montar_unidades(dados) -> None:
    secoes, votos, nomes = dados
    u = tr.montar_unidades(secoes, votos, nomes, "secao")
    assert u.cat1 == ["ANA (P1)", "BIA (P2)", "CAIO (P3)", "DUDA (P4)", "Outros", "Branco/nulo", "Abstenção"]
    assert u.cat2 == ["ANA (P1)", "BIA (P2)", "Branco/nulo", "Abstenção"] and u.finalistas == u.cat2[:2]
    t = u.tabela
    # as categorias somam o eleitorado apto em cada turno (a abstenção fecha a conta)
    assert (t.select(pl.sum_horizontal(pl.col("^1:.*$"))).to_series() == t["APTOS_1"]).all()
    assert (t.select(pl.sum_horizontal(pl.col("^2:.*$"))).to_series() == t["APTOS_2"]).all()
    loc = tr.montar_unidades(secoes, votos, nomes, "local")
    assert loc.tabela.height == 120 and loc.tabela["APTOS_1"].sum() == t["APTOS_1"].sum()
    assert loc.tabela["1:Outros"].sum() == t["1:Outros"].sum()
    mun = tr.montar_unidades(secoes, votos, nomes, "municipio")
    assert mun.cat1 == ["ANA (P1)", "BIA (P2)", "Eliminados", "Branco/nulo", "Abstenção"]  # municípios: um grupo só
    so_um = tr.montar_unidades(secoes, votos, nomes, "local", municipio=1002)
    assert set(so_um.tabela["CD_MUNICIPIO"]) == {1002}


def test_grupo_sem_eleitores_sai_da_matriz(dados) -> None:
    """Achado em Petrópolis 2024: sem nenhum voto em "Outros", a linha dava NaN e a API caía."""
    secoes, votos, nomes = dados
    sem_e = votos.filter(pl.col("NR_VOTAVEL") != 50)
    u = tr.montar_unidades(secoes, sem_e, nomes, "local")
    assert "Outros" not in u.cat1 and "1:Outros" not in u.tabela.columns
    r = tr.resumo(tr.analisar(u, n_boot=10))
    assert all(np.isfinite(d["pct"]) for m in r["matriz"] for d in m["destinos"])


def test_segundo_turno_precisa_de_dois_candidatos(dados) -> None:
    secoes, votos, nomes = dados
    with pytest.raises(tr.v.TseDataError, match="2 candidatos"):
        tr.montar_unidades(secoes, votos.filter(~((pl.col("NR_TURNO") == 2) & (pl.col("NR_VOTAVEL") == 20))), nomes)


def test_prefeito_exige_municipio() -> None:
    with pytest.raises(ValueError, match="município"):
        tr.calcular(2024, "RJ", 11, "local", Path("nada"))


# --------------------------------------------------------------------------- análise
def test_analise_estratificada_recupera_cada_zona(dados) -> None:
    secoes, votos, nomes = dados
    u = tr.montar_unidades(secoes, votos, nomes, "secao")
    r = tr.analisar(u, n_boot=30)
    assert r.estrato == "zona" and r.estratos.height == 2
    x, _, w = u.matrizes()
    e, _ = tr.estratos(u, "zona")
    # a verdade da área = matrizes das zonas combinadas pelo eleitorado de cada categoria
    tot = np.stack([(x[e == z] * w[e == z, None]).sum(axis=0) for z in range(2)])
    verdade = tr._combinar(np.stack([B_ZONA[1], B_ZONA[2]]), tot)
    grandes = [i for i, c in enumerate(u.cat1) if c != "Outros"]      # "Outros" tem 0,4% do eleitorado
    assert np.abs(r.matriz - verdade)[grandes].max() < 0.05   # 480 seções: grupos de 5% erram até ~4 p.p.
    assert np.abs(r.matriz - verdade).max() < 0.10
    assert ((r.baixo <= r.matriz) & (r.matriz <= r.alto)).all()          # o IC contém a estimativa
    cobre = (r.baixo - 0.005 <= verdade) & (verdade <= r.alto + 0.005)
    assert cobre.mean() >= 0.75                                          # e a verdade, na maioria das células
    assert np.allclose(r.matriz.sum(axis=1), 1)
    # cada local recebe o destino dos eliminados da sua zona: CAIO foi mais para ANA na zona 2
    zonas = r.estratos.sort("ESTRATO")
    assert zonas["ELIM_PARA_ANA (P1)_PCT"][1] > zonas["ELIM_PARA_ANA (P1)_PCT"][0] + 10
    assert r.por_unidade["ESTRATO"].n_unique() == 2


def test_validacao_cruzada_bate_o_swing_uniforme(dados) -> None:
    secoes, votos, nomes = dados
    val = tr.analisar(tr.montar_unidades(secoes, votos, nomes, "secao"), n_boot=0).validacao
    assert val["rmse_modelo_medio_pp"] < val["rmse_matriz_unica_medio_pp"] < val["rmse_swing_medio_pp"]


def test_resumo_e_planilha(dados, tmp_path: Path) -> None:
    import openpyxl
    secoes, votos, nomes = dados
    res = tr.analisar(tr.montar_unidades(secoes, votos, nomes, "local"), n_boot=10)
    r = tr.resumo(res)
    assert [m["origem"] for m in r["matriz"]] == r["categorias_1t"]
    for m in r["matriz"]:
        assert abs(sum(d["pct"] for d in m["destinos"]) - 100) < 1e-6
        assert abs(sum(d["eleitores"] for d in m["destinos"]) - m["eleitores_1t"]) <= len(m["destinos"])
    ab = r["abstencao"]
    assert ab["extra_pp"] == pytest.approx(ab["pct_2t"] - ab["pct_1t"])
    assert "Abstenção" not in [n["origem"] for n in ab["novos_abstencionistas"]]
    assert r["residuos_a"]["acima"][0]["RESIDUO_A_PP"] >= r["residuos_a"]["abaixo"][0]["RESIDUO_A_PP"]
    tr.para_planilha(r, res.por_unidade, tmp_path / "t.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "t.xlsx")
    assert wb.sheetnames == ["Matriz", "Por unidade"] and wb["Por unidade"].max_row == 121


def test_mesma_semente_mesmo_resultado(dados) -> None:
    secoes, votos, nomes = dados
    u = tr.montar_unidades(secoes, votos, nomes, "local")
    a, b = tr.analisar(u, n_boot=15, validar_cv=False), tr.analisar(u, n_boot=15, validar_cv=False)
    assert np.array_equal(a.baixo, b.baixo) and np.array_equal(a.alto, b.alto)


# --------------------------------------------------------------------------- API
def test_api(dados, tse_cache: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    secoes, votos, nomes = dados
    chamadas = []

    def calcular(ano, uf, cargo, nivel, cache, municipio=None, n_boot=200):
        chamadas.append((ano, cargo, nivel, municipio))
        return tr.analisar(tr.montar_unidades(secoes, votos, nomes, nivel, municipio), n_boot=5)

    monkeypatch.setattr(tr, "calcular", calcular)
    site = TestClient(create_app(tmp_path / "dados", "RJ", tse_cache))
    info = site.get("/api/transferencia/info").json()
    assert info["tempo_real"] is False and 2024 in info["anos"] and info["cargos"]["2024"] == [11]
    r = site.get("/api/transferencia?ano=2022&cargo=1&nivel=secao").json()
    assert r["unidades"] == 480 and r["n_estratos"] == 2
    site.get("/api/transferencia?ano=2022&cargo=1&nivel=secao")
    assert len(chamadas) == 1                                        # 2º pedido: do cache
    assert site.get("/api/transferencia?fonte=tempo_real&cargo=1").status_code == 400   # site do 1º turno
    assert site.get("/api/transferencia?nivel=bairro").status_code == 400
    na_rede = TestClient(create_app(tmp_path / "dados", "RJ", tse_cache, pesadas_so_local=True),
                         client=("192.168.0.20", 5000))
    assert na_rede.get("/api/transferencia?ano=2022&cargo=1").status_code == 403


def test_cli(dados, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture) -> None:
    import transferencia_turnos as cli
    monkeypatch.setattr(tr, "carregar_microdados", lambda ano, uf, cargo, cache, municipio=None: dados)
    assert cli.main(["--nivel", "local", "--bootstrap", "5", "--saida", str(tmp_path / "t.xlsx")]) == 0
    saida = capsys.readouterr().out
    assert "Para onde foi cada grupo" in saida and "swing uniforme" in saida and (tmp_path / "t.xlsx").exists()
    assert cli.main(["--comparar-niveis"]) == 0
    assert "maior diferença para as seções" in capsys.readouterr().out


# --------------------------------------------------------------------------- 2022 real (regressão)
CACHE_REAL = Path(__file__).parent / "cache_tse"
HIST = Path(__file__).parent / "dados_2026"
tem_2022 = pytest.mark.skipif(not (CACHE_REAL / "votacao_secao_2022_BR.zip").exists()
                              or not (CACHE_REAL / "detalhe_votacao_secao_2022.zip").exists(),
                              reason="microdados de 2022 fora do cache")


@tem_2022
def test_presidente_2022_rj_por_local() -> None:
    r = tr.calcular(2022, "RJ", 1, "local", CACHE_REAL, n_boot=20)
    cat = r.unidades.cat1
    assert cat[:2] == ["JAIR BOLSONARO (PL)", "LULA (PT)"] and r.unidades.tabela.height > 4500
    m = 100 * r.matriz
    assert m[0, 0] > 95 and m[1, 1] > 90                             # quem votou num finalista voltou nele
    tebet = cat.index("SIMONE TEBET (MDB)")
    assert 40 < m[tebet, 1] < 55 and 20 < m[tebet, 0] < 35
    val = r.validacao
    assert val["rmse_modelo_medio_pp"] < val["rmse_matriz_unica_medio_pp"] < val["rmse_swing_medio_pp"]


@tem_2022
@pytest.mark.skipif(not (HIST / "historico_2022_t2" / "ultimo").exists(), reason="2022 não importado")
def test_tempo_real_igual_aos_microdados_por_municipio() -> None:
    """O caminho da noite (o que o coletor grava nos dois turnos) dá os mesmos números que os microdados."""
    u = tr.unidades_divulgacao(HIST / "historico_2022_t1", HIST / "historico_2022_t2", 1, "RJ")
    secoes, votos, nomes = tr.carregar_microdados(2022, "RJ", 1, CACHE_REAL)
    um = tr.montar_unidades(secoes, votos, nomes, "municipio")
    assert u.cat1 == um.cat1
    a, b = u.tabela.sort("UNIDADE"), um.tabela.sort("UNIDADE")
    assert a.select(pl.exclude("NOME", "NM_MUNICIPIO", "GRUPO")).equals(b.select(pl.exclude("NOME", "NM_MUNICIPIO", "GRUPO")))
